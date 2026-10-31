"""
============================================================
IMPORTS
============================================================
"""
import os
import io
import re
import traceback
from collections import namedtuple, defaultdict
from datetime import datetime

import pandas as pd
from flask import (
    Flask, render_template, request, redirect, url_for,
    session, flash, send_file, jsonify,
)
from sqlalchemy import or_, func
from sqlalchemy.orm import joinedload
from sqlalchemy.exc import IntegrityError
from werkzeug.utils import secure_filename
from openpyxl.styles import Font, PatternFill, Alignment

from connections import SessionLocal
from models import (
    User, Admin, HOD, Department, Course, Semester, Subject,
    Question, Quiz, StudentProfile, Result, Message,
    ActivityLog, KnecMark, Module,
    ShortCourseMark, PeriodicMark,
)
from utils import parse_docx_questions


"""
============================================================
APP CONFIG
============================================================
"""
app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 100 * 1024 * 1024
app.secret_key = "132silas456sakwa789ayanga"


"""
============================================================
QUERY HELPERS
============================================================
"""
def student_query(db):
    return db.query(StudentProfile).options(
        joinedload(StudentProfile.course),
        joinedload(StudentProfile.current_semester),
    )


def subject_query(db):
    return db.query(Subject).options(
        joinedload(Subject.course),
        joinedload(Subject.semester),
        joinedload(Subject.module_obj),
    )


def semester_query(db):
    return db.query(Semester).options(joinedload(Semester.course))


def quiz_query(db):
    return db.query(Quiz).options(
        joinedload(Quiz.course),
        joinedload(Quiz.subject),
    )


"""
============================================================
MANUAL PAGINATION (SQLAlchemy 2.x compatible)
============================================================
"""
class ManualPagination:
    def __init__(self, items, page, per_page, total):
        self.items = items
        self.page = page
        self.per_page = per_page
        self.total = total
        self.pages = (total + per_page - 1) // per_page if per_page else 0

    @property
    def has_prev(self):
        return self.page > 1

    @property
    def has_next(self):
        return self.page < self.pages

    @property
    def prev_num(self):
        return self.page - 1 if self.has_prev else None

    @property
    def next_num(self):
        return self.page + 1 if self.has_next else None

    @property
    def first(self):
        if not self.total:
            return 0
        return (self.page - 1) * self.per_page + 1

    @property
    def last(self):
        if not self.total:
            return 0
        return min(self.page * self.per_page, self.total)

    def iter_pages(self, left_edge=1, right_edge=1, left_current=2, right_current=2):
        last = 0
        for num in range(1, self.pages + 1):
            if (
                num <= left_edge
                or (num > self.page - left_current - 1 and num < self.page + right_current)
                or num > self.pages - right_edge
            ):
                if last + 1 != num:
                    yield None
                yield num
                last = num


def paginate_query(query, page, per_page):
    if page < 1:
        page = 1
    total = query.order_by(None).count()
    items = query.offset((page - 1) * per_page).limit(per_page).all()
    return ManualPagination(items=items, page=page, per_page=per_page, total=total)


"""
============================================================
ASSESSMENT TYPE HELPERS
============================================================
"""
VALID_ASSESSMENT_TYPES = ("knec", "kasneb", "short", "periodic")
KASNEB_LABELS = ("CAT 1", "CAT 2")


def get_course_assessment_type(course):
    return (getattr(course, "assessment_type", None) or "knec").strip().lower()


def get_periodic_labels(course):
    if get_course_assessment_type(course) != "periodic":
        return []
    raw = (getattr(course, "periodic_labels", None) or "").strip()
    if not raw:
        return ["Month 1", "Month 2", "Month 3", "Final"]
    return [l.strip() for l in raw.split(",") if l.strip()]


def get_student_subject_marks(db, student_id, subject_id, semester_id, course):
    atype = get_course_assessment_type(course)

    if atype == "knec":
        m = db.query(KnecMark).filter_by(
            student_id=student_id, subject_id=subject_id, semester_id=semester_id
        ).first()
        if not m:
            return None
        return {
            "cat1": m.cat1, "cat2": m.cat2, "final": m.final,
            "total": (m.cat1 or 0) + (m.cat2 or 0) + (m.final or 0),
        }

    if atype == "short":
        m = db.query(ShortCourseMark).filter_by(
            student_id=student_id, subject_id=subject_id, semester_id=semester_id
        ).first()
        if not m:
            return None
        return {"score": m.score, "remarks": m.remarks, "total": m.score or 0}

    if atype == "kasneb":
        rows = db.query(PeriodicMark).filter_by(
            student_id=student_id, subject_id=subject_id, semester_id=semester_id
        ).all()

        cells = {r.label: r.score for r in rows}
        cat1 = cells.get("CAT 1")
        cat2 = cells.get("CAT 2")

        if cat1 is not None or cat2 is not None:
            scores = [c for c in (cat1, cat2) if c is not None]
            avg = round(sum(scores) / len(scores), 2) if scores else None
            return {
                "cells": cells,
                "labels": list(KASNEB_LABELS),
                "average": avg,
                "total": avg or 0,
            }

        quiz_results = db.query(Result).join(Quiz).filter(
            Result.student_id == student_id,
            Quiz.subject_id == subject_id
        ).all()

        if quiz_results:
            best_score = max([r.percentage for r in quiz_results])
            return {
                "cells": {"CAT 1": best_score, "CAT 2": None},
                "labels": list(KASNEB_LABELS),
                "average": round(best_score, 2),
                "total": best_score,
            }

        return None

    labels = get_periodic_labels(course)
    rows = db.query(PeriodicMark).filter_by(
        student_id=student_id, subject_id=subject_id, semester_id=semester_id
    ).all()
    if not rows:
        return None
    cells = {r.label: r.score for r in rows}
    scores = [c for c in cells.values() if c is not None]
    avg = round(sum(scores) / len(scores), 2) if scores else None
    return {"cells": cells, "labels": labels, "average": avg, "total": avg or 0}


def knec_grade(avg):
    if avg is None:
        return ""
    return (1 if avg >= 84 else 2 if avg >= 75 else 3 if avg >= 66
            else 4 if avg >= 60 else 5 if avg >= 50
            else 6 if avg >= 40 else 7)


def knec_result(avg):
    if avg is None:
        return ""
    return ("DISTINCTION" if avg >= 75 else "CREDIT" if avg >= 60
            else "PASS" if avg >= 40 else "REFER")


"""
============================================================
KNEC ANALYTICS WORKBOOK BUILDER (OPTIMIZED)
============================================================
"""
def _build_knec_analytics_workbook(db, course):
    students = (
        student_query(db).filter_by(course_id=course.id)
        .order_by(StudentProfile.full_name.asc()).all()
    )
    semesters = (
        semester_query(db).filter_by(course_id=course.id)
        .order_by(Semester.created_at.asc()).all()
    )

    all_subjects = (
        subject_query(db).filter_by(course_id=course.id)
        .order_by(Subject.name.asc()).all()
    )
    subject_by_id = {s.id: s for s in all_subjects}

    student_ids = [s.id for s in students]
    all_marks = []
    if student_ids:
        all_marks = (
            db.query(KnecMark)
            .filter(KnecMark.student_id.in_(student_ids))
            .all()
        )

    marks_by_student_sem = {}
    for m in all_marks:
        sem_obj = next((s for s in semesters if s.id == m.semester_id), None)
        if not sem_obj:
            continue
        sem_name = sem_obj.name
        marks_by_student_sem.setdefault(m.student_id, {}).setdefault(sem_name, {})[m.subject_id] = m

    student_rows = []
    raw_rows = []
    subject_stats = {}

    for s in students:
        per_sem_avg = {}
        overall_sum, overall_n = 0, 0

        for sem in semesters:
            sem_marks = marks_by_student_sem.get(s.id, {}).get(sem.name, {})
            sem_sum, sem_n = 0, 0

            for subj_id, m in sem_marks.items():
                subj = subject_by_id.get(subj_id)
                subj_name = subj.name if subj else f"Subject #{subj_id}"
                total = (m.cat1 or 0) + (m.cat2 or 0) + (m.final or 0)
                sem_sum += total
                sem_n += 1
                overall_sum += total
                overall_n += 1
                raw_rows.append({
                    "Admission No": s.admission_number,
                    "Student Name": s.full_name,
                    "Semester": sem.name,
                    "Subject": subj_name,
                    "CAT1 (15)": m.cat1,
                    "CAT2 (15)": m.cat2,
                    "Final (70)": m.final,
                    "Total (100)": total,
                })
                subject_stats.setdefault(subj_id, []).append(total)

            per_sem_avg[sem.name] = round(sem_sum / sem_n, 2) if sem_n else None

        overall_avg = round(overall_sum / overall_n, 2) if overall_n else None
        row = {
            "Admission No": s.admission_number,
            "Student Name": s.full_name,
            "Exam Type": (s.exam_type or "").upper(),
        }
        row.update(per_sem_avg)
        row["Overall Average"] = overall_avg
        row["Grade"] = knec_grade(overall_avg)
        row["Result"] = knec_result(overall_avg)
        student_rows.append(row)

    stat_rows = []
    for subj_id, totals in subject_stats.items():
        subj = subject_by_id.get(subj_id)
        n = len(totals)
        mean = round(sum(totals) / n, 2) if n else 0
        variance = round(sum((x - mean) ** 2 for x in totals) / n, 2) if n else 0
        stat_rows.append({
            "Subject": subj.name if subj else "",
            "Entries": n,
            "Mean": mean,
            "Std Dev": round(variance ** 0.5, 2),
            "Highest": max(totals) if totals else 0,
            "Lowest": min(totals) if totals else 0,
            "Pass Rate (>=40)": round(sum(1 for x in totals if x >= 40) / n * 100, 2) if n else 0,
            "Distinction Rate (>=75)": round(sum(1 for x in totals if x >= 75) / n * 100, 2) if n else 0,
        })
    stat_rows.sort(key=lambda r: r["Subject"])

    dist = defaultdict(int)
    for r in student_rows:
        if r["Result"]:
            dist[r["Result"]] += 1
    dist_rows = [{"Result": k, "Students": v} for k, v in dist.items()]

    ranked = sorted(
        [r for r in student_rows if r["Overall Average"] is not None],
        key=lambda r: r["Overall Average"], reverse=True,
    )
    rank_rows = []
    for i, r in enumerate(ranked, 1):
        rank_rows.append({
            "Rank": i,
            "Admission No": r["Admission No"],
            "Student Name": r["Student Name"],
            "Overall Average": r["Overall Average"],
            "Grade": r["Grade"],
            "Result": r["Result"],
        })

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        sheets = {
            "Student Summary": pd.DataFrame(student_rows),
            "Ranking": pd.DataFrame(rank_rows),
            "Subject Stats": pd.DataFrame(stat_rows),
            "Grade Distribution": pd.DataFrame(dist_rows),
            "Raw Marks": pd.DataFrame(raw_rows),
        }
        for name, df in sheets.items():
            if df.empty:
                df = pd.DataFrame({"Note": ["No data available"]})
            df.to_excel(writer, index=False, sheet_name=name)
            ws = writer.sheets[name]

            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill("solid", fgColor="1ABC9C")
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for col in ws.columns:
                width = max(len(str(c.value or "")) for c in col) + 3
                ws.column_dimensions[col[0].column_letter].width = min(width, 40)

            ws.freeze_panes = "A2"

    buf.seek(0)
    return buf


"""
============================================================
CONTEXT PROCESSORS
============================================================
"""
@app.context_processor
def inject_hod_context():
    if session.get("role") == "hod":
        db = SessionLocal()
        try:
            dept = db.query(Department).filter_by(id=session.get("department_id")).first()
            return {"department": dept, "hod_name": session.get("username")}
        finally:
            db.close()
    return {}


@app.context_processor
def inject_semester_helpers():
    from datetime import datetime as _dt

    MONTH_MAP = {
        'jan': 1, 'january': 1, 'feb': 2, 'february': 2,
        'mar': 3, 'march': 3, 'apr': 4, 'april': 4,
        'may': 5, 'jun': 6, 'june': 6, 'jul': 7, 'july': 7,
        'aug': 8, 'august': 8, 'sep': 9, 'sept': 9, 'september': 9,
        'oct': 10, 'october': 10, 'nov': 11, 'november': 11,
        'dec': 12, 'december': 12,
    }

    def parse_semester_months(name):
        if not name:
            return None, None
        parts = re.split(r'\s*(?:–|—|-|/|to|&)\s*', name.lower())
        months = []
        for part in parts:
            word = part.strip().split()[0] if part.strip() else ''
            if word in MONTH_MAP:
                months.append(MONTH_MAP[word])
        if len(months) >= 2:
            return months[0], months[-1]
        elif len(months) == 1:
            return months[0], months[0]
        return None, None

    def semester_time_status(semester):
        now_month = _dt.now().month
        start, end = parse_semester_months(semester.name if semester else '')

        if semester is None:
            return {"state": "pending", "label": "Pending", "short": ""}

        if start is None:
            if getattr(semester, 'is_current', False):
                return {"state": "current", "label": "Current", "short": semester.name}
            return {"state": "pending", "label": "Pending", "short": semester.name}

        if start <= end:
            in_range = start <= now_month <= end
            before = now_month < start
        else:
            in_range = now_month >= start or now_month <= end
            before = end < now_month < start

        if in_range or getattr(semester, 'is_current', False):
            return {"state": "current", "label": "Current", "short": semester.name}
        elif before:
            return {"state": "next", "label": "Next", "short": semester.name}
        else:
            return {"state": "done", "label": "Done", "short": semester.name}

    return {"semester_time_status": semester_time_status}


"""
============================================================
UPLOAD FOLDERS
============================================================
"""
BASE_UPLOAD = os.path.join(os.getcwd(), "static", "uploads")
EXAMS_UPLOAD_FOLDER = os.path.join(BASE_UPLOAD, "exams")
QUESTION_IMAGES_FOLDER = os.path.join("static", "question_images")

for path in [BASE_UPLOAD, EXAMS_UPLOAD_FOLDER, QUESTION_IMAGES_FOLDER]:
    os.makedirs(path, exist_ok=True)

ALLOWED_EXAM_EXTENSIONS = {"docx"}


def allowed(filename, allowed_set):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in allowed_set


"""
============================================================
SHARED HELPERS
============================================================
"""
def _delete_quiz_tree(db, quiz):
    db.query(Result).filter(Result.quiz_id == quiz.id).delete(synchronize_session=False)
    db.query(Question).filter(Question.quiz_id == quiz.id).delete(synchronize_session=False)
    db.delete(quiz)


def _delete_quizzes_for(db, course_id=None, subject_id=None):
    q = db.query(Quiz)
    if course_id is not None:
        q = q.filter(Quiz.course_id == course_id)
    if subject_id is not None:
        q = q.filter(Quiz.subject_id == subject_id)
    for quiz in q.all():
        _delete_quiz_tree(db, quiz)


def _delete_student_tree(db, student):
    db.query(ActivityLog).filter(ActivityLog.student_id == student.id).delete(synchronize_session=False)
    db.query(Result).filter(Result.student_id == student.id).delete(synchronize_session=False)
    db.query(KnecMark).filter(KnecMark.student_id == student.id).delete(synchronize_session=False)
    db.query(ShortCourseMark).filter(ShortCourseMark.student_id == student.id).delete(synchronize_session=False)
    db.query(PeriodicMark).filter(PeriodicMark.student_id == student.id).delete(synchronize_session=False)
    user = student.user
    db.delete(student)
    if user:
        db.delete(user)


def is_admin():
    return session.get("role") == "admin"


def is_hod():
    return session.get("role") == "hod"


def require_admin_or_hod():
    if session.get("role") not in ("admin", "hod"):
        return redirect(url_for("login"))
    return None


def require_admin():
    if session.get("role") != "admin":
        flash("Only the admin can manage courses and semesters.", "danger")
        return redirect(url_for("login"))
    return None


def _save_kasneb_marks(db, student_id, subject_id, semester_id, cat1, cat2):
    for label, value in (("CAT 1", cat1), ("CAT 2", cat2)):
        existing = db.query(PeriodicMark).filter_by(
            student_id=student_id, subject_id=subject_id,
            semester_id=semester_id, label=label,
        ).first()

        if value is None:
            if existing:
                db.delete(existing)
            continue

        if existing:
            existing.score = value
            existing.entered_at = datetime.utcnow()
        else:
            db.add(PeriodicMark(
                student_id=student_id, subject_id=subject_id,
                semester_id=semester_id, label=label, score=value,
            ))
    return 1


def _create_single_student(db, *, full_name, admission, course, exam_type,
                           semester_id, module_obj, phone, username, password,
                           strict_department_id=None):
    if db.query(StudentProfile).filter_by(admission_number=admission).first():
        return False, f"Admission number '{admission}' already exists.", None

    if strict_department_id is not None and course.department_id != strict_department_id:
        return False, "Course not in your department.", None

    if get_course_assessment_type(course) in ("short", "periodic"):
        exam_type = "internal"

    if semester_id is not None:
        sem = db.query(Semester).filter_by(id=semester_id, course_id=course.id).first()
        if not sem:
            return False, "Selected semester does not belong to the chosen course.", None
    else:
        cur = db.query(Semester).filter_by(
            course_id=course.id, is_current=True
        ).first()
        semester_id = cur.id if cur else None

    if module_obj is not None and module_obj.course_id != course.id:
        return False, "Selected module does not belong to the chosen course.", None

    base_username = (username or admission.replace("/", "").replace(" ", "").lower())[:45]
    if not base_username:
        base_username = admission.replace("/", "").replace(" ", "").lower()[:45]
    final_username = base_username
    counter = 1
    while db.query(User).filter_by(username=final_username).first():
        final_username = f"{base_username}{counter}"
        counter += 1

    user = User(username=final_username, password=password or admission)
    db.add(user)
    db.flush()

    db.add(StudentProfile(
        user_id=user.id,
        full_name=full_name,
        admission_number=admission,
        course_id=course.id,
        exam_type=exam_type,
        phone_number=phone,
        current_semester_id=semester_id,
        module=module_obj.name if module_obj else None,
        module_id=module_obj.id if module_obj else None,
    ))

    return True, f"Student '{full_name}' added (username: {final_username}).", final_username


def _validate_new_student_form(db, form, *, department_id=None):
    errors = []

    full_name = (form.get("full_name") or "").strip()
    admission = (form.get("admission_number") or "").strip()
    course_id = form.get("course_id", type=int)
    exam_type = (form.get("exam_type") or "").strip().lower()
    module_name = (form.get("module") or "").strip()
    semester_id = form.get("semester_id", type=int)
    phone = (form.get("phone_number") or "").strip()
    username = (form.get("username") or "").strip()
    password = (form.get("password") or "").strip()

    if not full_name:
        errors.append("Full name is required.")
    elif len(full_name) < 3:
        errors.append("Full name must be at least 3 characters.")

    if not admission:
        errors.append("Admission number is required.")
    else:
        existing = db.query(StudentProfile).filter_by(admission_number=admission).first()
        if existing:
            course_label = existing.course.name if existing.course else "another course"
            errors.append(
                f"Admission number '{admission}' is already used by "
                f"{existing.full_name} ({course_label})."
            )

    course = None
    if not course_id:
        errors.append("Course is required.")
    else:
        course = db.query(Course).filter_by(id=course_id).first()
        if not course:
            errors.append(f"Course #{course_id} does not exist.")
        elif department_id is not None and course.department_id != department_id:
            errors.append(f"'{course.name}' is not in your department.")

    if exam_type not in ("knec", "internal"):
        errors.append("Exam type must be KNEC or Internal.")

    module_obj = None
    if module_name and course:
        module_obj = db.query(Module).filter_by(
            course_id=course.id, name=module_name
        ).first()
        if not module_obj:
            errors.append(
                f"Module '{module_name}' not found for course '{course.name}'."
            )

    if semester_id:
        sem = db.query(Semester).filter_by(id=semester_id).first()
        if not sem:
            errors.append(f"Semester #{semester_id} does not exist.")
        elif course and sem.course_id != course.id:
            errors.append(
                f"Semester '{sem.name}' does not belong to course '{course.name}'."
            )

    if username and db.query(User).filter_by(username=username).first():
        errors.append(f"Username '{username}' is already taken.")

    if password and len(password) < 4:
        errors.append("Password must be at least 4 characters.")

    return {
        "full_name": full_name,
        "admission": admission,
        "course": course,
        "exam_type": exam_type,
        "module_obj": module_obj,
        "module_name": module_name,
        "semester_id": semester_id,
        "phone": phone or None,
        "username": username,
        "password": password,
    }, errors


def _error_fields_from(errors):
    fields = set()
    for e in errors:
        e_low = e.lower()
        if "full name" in e_low:
            fields.add("full_name")
        elif "admission" in e_low:
            fields.add("admission_number")
        elif "course" in e_low:
            fields.add("course_id")
        elif "exam type" in e_low:
            fields.add("exam_type")
        elif "module" in e_low:
            fields.add("module")
        elif "semester" in e_low:
            fields.add("semester_id")
        elif "phone" in e_low:
            fields.add("phone_number")
        elif "username" in e_low:
            fields.add("username")
        elif "password" in e_low:
            fields.add("password")
    return fields


def _translate_db_error(exc):
    err = str(getattr(exc, "orig", exc)).lower()
    if "unique" in err and "admission_number" in err:
        return "That admission number is already in use."
    if "unique" in err and "username" in err:
        return "That username is already taken. Try a different one."
    if "not null" in err:
        missing = err.split("failed:")[-1].strip() if "failed:" in err else "a required field"
        return f"Required field is missing: {missing}"
    if "foreign key" in err:
        return "One of the selected values (course / semester / module) no longer exists."
    return "Could not save the record. Please check the form and try again."


"""
============================================================
PUBLIC ROUTES
============================================================
"""
@app.route("/")
def home():
    return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    error = False
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")

        db = SessionLocal()
        try:
            admin = db.query(Admin).filter_by(username=username).first()
            if admin and admin.password == password:
                session.update({"username": admin.username, "user_id": admin.id, "role": "admin"})
                return redirect(url_for("admin_dashboard"))

            hod = db.query(HOD).filter_by(username=username).first()
            if hod and hod.password == password:
                session.update({
                    "username": hod.username,
                    "user_id": hod.id,
                    "role": "hod",
                    "department_id": hod.department_id,
                })
                return redirect(url_for("hod_dashboard"))

            user = db.query(User).filter_by(username=username).first()
            if user and user.password == password:
                session.update({"username": user.username, "user_id": user.id, "role": "student"})
                return redirect(url_for("student_dashboard"))

            error = True
        finally:
            db.close()

    return render_template("login.html", error=error)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        username = request.form["username"].strip()
        password = request.form["password"]

        db = SessionLocal()
        try:
            if db.query(User).filter_by(username=username).first():
                flash("❌ Username already exists. Please choose another one.", "danger")
                return render_template("students/Register.html", username=username)

            db.add(User(username=username, password=password))
            db.commit()
            flash("✅ Registration successful! Please log in.", "success")
            return redirect(url_for("login"))
        except Exception:
            db.rollback()
            flash("❌ Something went wrong. Try again.", "danger")
            return render_template("students/Register.html", username=username)
        finally:
            db.close()

    return render_template("students/Register.html")


"""
============================================================
JSON HELPERS (cascading dropdowns)
============================================================
"""
@app.route("/api/semesters_by_course/<int:course_id>")
def api_semesters_by_course(course_id):
    if session.get("role") not in ("admin", "hod"):
        return jsonify([]), 403
    db = SessionLocal()
    try:
        rows = db.query(Semester).filter_by(course_id=course_id)\
               .order_by(Semester.created_at.asc()).all()
        return jsonify([{"id": s.id, "name": s.name} for s in rows])
    finally:
        db.close()


@app.route("/api/modules_by_course/<int:course_id>")
def api_modules_by_course(course_id):
    if session.get("role") not in ("admin", "hod"):
        return jsonify([]), 403
    db = SessionLocal()
    try:
        rows = db.query(Module).filter_by(course_id=course_id)\
               .order_by(Module.order.asc(), Module.name.asc()).all()
        return jsonify([{"id": m.id, "name": m.name} for m in rows])
    finally:
        db.close()


"""
============================================================
ADMIN ROUTES
============================================================
"""
@app.route("/admin_dashboard")
def admin_dashboard():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        courses = db.query(Course).all()
        students_count = db.query(StudentProfile).count()
        exams_count = db.query(Quiz).count()

        departments_list = db.query(Department).order_by(Department.name.asc()).all()
        departments = []
        for d in departments_list:
            dept_courses = db.query(Course).filter_by(department_id=d.id).all()
            dept_course_ids = [c.id for c in dept_courses]
            dept_students = 0
            if dept_course_ids:
                dept_students = db.query(StudentProfile).filter(StudentProfile.course_id.in_(dept_course_ids)).count()
            departments.append({
                "id": d.id,
                "name": d.name,
                "courses_count": len(dept_courses),
                "students_count": dept_students
            })
        departments_count = len(departments)

        return render_template(
            "admin/admin_dashboard.html",
            courses=courses,
            departments=departments,
            students_count=students_count,
            exams_count=exams_count,
            departments_count=departments_count,
            username=session.get("username"),
            current_year=datetime.now().year,
        )
    finally:
        db.close()


@app.route("/admin/departments", methods=["GET", "POST"])
def manage_departments():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        if request.method == "POST":
            name = request.form.get("name", "").strip()
            code = request.form.get("code", "").strip() or None

            if not name:
                flash("Department name is required.", "danger")
                return redirect(url_for("manage_departments"))

            if db.query(Department).filter_by(name=name).first():
                flash(f"Department '{name}' already exists.", "warning")
                return redirect(url_for("manage_departments"))

            db.add(Department(name=name, code=code))
            db.commit()
            flash(f"✅ Department '{name}' added.", "success")
            return redirect(url_for("manage_departments"))

        departments = db.query(Department).order_by(Department.name.asc()).all()
        dept_stats = []
        for d in departments:
            dept_stats.append({
                "department": d,
                "courses": db.query(Course).filter_by(department_id=d.id).count(),
                "hods": db.query(HOD).filter_by(department_id=d.id).count(),
            })

        return render_template("admin/manage_departments.html", dept_stats=dept_stats)
    finally:
        db.close()


@app.route("/admin/departments/delete/<int:dept_id>", methods=["POST"])
def delete_department(dept_id):
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        dept = db.query(Department).filter_by(id=dept_id).first()
        if not dept:
            flash("Department not found.", "danger")
            return redirect(url_for("manage_departments"))

        if db.query(Course).filter_by(department_id=dept_id).count():
            flash("Cannot delete — courses still in this department.", "danger")
            return redirect(url_for("manage_departments"))

        if db.query(HOD).filter_by(department_id=dept_id).count():
            flash("Cannot delete — an HOD is still assigned.", "danger")
            return redirect(url_for("manage_departments"))

        db.delete(dept)
        db.commit()
        flash("✅ Department deleted.", "success")
    finally:
        db.close()

    return redirect(url_for("manage_departments"))


@app.route("/admin/hods", methods=["GET", "POST"])
def manage_hods():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        if request.method == "POST":
            username = request.form.get("username", "").strip()
            password = request.form.get("password", "")
            full_name = request.form.get("full_name", "").strip()
            department_id = request.form.get("department_id")

            if not all([username, password, full_name, department_id]):
                flash("All fields are required.", "danger")
                return redirect(url_for("manage_hods"))

            if db.query(HOD).filter_by(username=username).first():
                flash(f"Username '{username}' already exists.", "warning")
                return redirect(url_for("manage_hods"))

            if db.query(HOD).filter_by(department_id=int(department_id)).first():
                flash("That department already has an HOD.", "warning")
                return redirect(url_for("manage_hods"))

            db.add(HOD(
                username=username, password=password,
                full_name=full_name, department_id=int(department_id),
            ))
            db.commit()
            flash(f"✅ HOD '{username}' added.", "success")
            return redirect(url_for("manage_hods"))

        departments = db.query(Department).order_by(Department.name.asc()).all()
        hods = (
            db.query(HOD)
            .options(joinedload(HOD.department))
            .order_by(HOD.username.asc())
            .all()
        )
        taken = {h.department_id for h in hods}

        return render_template(
            "admin/manage_hods.html",
            departments=departments, hods=hods, taken_departments=taken,
        )
    finally:
        db.close()


@app.route("/admin/hods/delete/<int:hod_id>", methods=["POST"])
def delete_hod(hod_id):
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        hod = db.query(HOD).filter_by(id=hod_id).first()
        if hod:
            db.delete(hod)
            db.commit()
            flash("✅ HOD deleted.", "success")
        else:
            flash("HOD not found.", "danger")
    finally:
        db.close()

    return redirect(url_for("manage_hods"))


@app.route("/manage_courses", methods=["GET"])
def manage_courses():
    blocked = require_admin_or_hod()
    if blocked:
        return blocked

    db = SessionLocal()
    try:
        if is_hod():
            base_q = db.query(Course).filter_by(department_id=session["department_id"])
        else:
            base_q = db.query(Course)

        search_query = request.args.get("search", "").strip()
        if search_query:
            like = f"%{search_query}%"
            base_q = base_q.filter(or_(
                Course.name.ilike(like),
                Course.level.ilike(like)
            ))

        sem_filter = request.args.get("semester_filter", "").strip()
        if sem_filter == "missing":
            subq = db.query(Semester.course_id).distinct()
            base_q = base_q.filter(~Course.id.in_(subq))
        elif sem_filter == "has":
            subq = db.query(Semester.course_id).distinct()
            base_q = base_q.filter(Course.id.in_(subq))

        page = request.args.get("page", 1, type=int)
        per_page = 10
        pagination = paginate_query(base_q.order_by(Course.name.asc()), page, per_page)
        courses = pagination.items

        for c in courses:
            c.students_count = db.query(StudentProfile).filter_by(course_id=c.id).count()
            c.subjects_count = db.query(Subject).filter_by(course_id=c.id).count()

        return render_template(
            "admin/manage_courses.html",
            courses=courses,
            pagination=pagination,
            search_query=search_query,
            sem_filter=sem_filter,
            hod_mode=is_hod(),
        )
    finally:
        db.close()


@app.route("/add_course", methods=["GET", "POST"])
def add_course():
    blocked = require_admin()
    if blocked:
        return blocked

    db = SessionLocal()
    try:
        if request.method == "POST":
            course_name = request.form["course_name"].strip()
            course_level = request.form["course_level"].strip()
            department_id = request.form.get("department_id") or None
            department_id = int(department_id) if department_id else None

            dept = None
            if department_id:
                dept = db.query(Department).filter_by(id=department_id).first()

            if dept and dept.name.strip().upper() in ("CPA", "KASNEB"):
                assessment_type = "short"
                periodic_labels = None
                max_score = 100
            else:
                assessment_type = (request.form.get("assessment_type") or "knec").strip().lower()
                if assessment_type not in ("knec", "short", "periodic"):
                    assessment_type = "knec"

                max_score = request.form.get("max_score", type=int) or 100
                if max_score <= 0 or max_score > 1000:
                    max_score = 100

                periodic_labels = (request.form.get("periodic_labels") or "").strip()
                if assessment_type == "periodic" and not periodic_labels:
                    periodic_labels = "Month 1,Month 2,Month 3,Final"
                if assessment_type != "periodic":
                    periodic_labels = None

            if db.query(Course).filter_by(name=course_name).first():
                flash(f"Course '{course_name}' already exists!", "warning")
                return redirect(url_for("add_course"))

            db.add(Course(
                name=course_name, level=course_level,
                department_id=department_id,
                assessment_type=assessment_type,
                max_score=max_score,
                periodic_labels=periodic_labels,
            ))
            db.commit()
            flash(f"✅ Course added ({assessment_type.upper()}).", "success")
            return redirect(url_for("manage_courses"))

        departments = db.query(Department).order_by(Department.name.asc()).all()
        return render_template("admin/add_course.html", departments=departments)
    except Exception as e:
        db.rollback()
        print("❌ Error in /add_course:", e)
        flash("An error occurred while adding the course.", "danger")
        return redirect(url_for("add_course"))
    finally:
        db.close()


@app.route("/edit_course/<int:course_id>", methods=["GET", "POST"])
def edit_course(course_id):
    blocked = require_admin()
    if blocked:
        return blocked

    db = SessionLocal()
    try:
        course = db.query(Course).filter(Course.id == course_id).first()
        if not course:
            return "Course not found", 404

        if request.method == "POST":
            course.name = request.form["course_name"].strip()
            course.level = request.form["course_level"].strip()
            dept = request.form.get("department_id") or None
            course.department_id = int(dept) if dept else None

            dept_obj = None
            if course.department_id:
                dept_obj = db.query(Department).filter_by(id=course.department_id).first()

            if dept_obj and dept_obj.name.strip().upper() in ("CPA", "KASNEB"):
                course.assessment_type = "short"
                course.periodic_labels = None
                course.max_score = 100
            else:
                atype = (request.form.get("assessment_type") or course.assessment_type or "knec").strip().lower()
                if atype not in ("knec", "short", "periodic"):
                    atype = "knec"
                course.assessment_type = atype
                course.max_score = request.form.get("max_score", type=int) or course.max_score or 100
                labels = (request.form.get("periodic_labels") or "").strip()
                course.periodic_labels = labels if atype == "periodic" else None

            db.commit()
            flash("Course updated.", "success")
            return redirect(url_for("manage_courses"))

        departments = db.query(Department).order_by(Department.name.asc()).all()
        return render_template("admin/edit_course.html", course=course, departments=departments)
    finally:
        db.close()


@app.route("/delete_course/<int:course_id>", methods=["POST"])
def delete_course(course_id):
    blocked = require_admin()
    if blocked:
        return blocked

    db = SessionLocal()
    try:
        course = db.query(Course).filter_by(id=course_id).first()
        if not course:
            return "Course not found", 404

        enrolled = db.query(StudentProfile).filter_by(course_id=course_id).count()
        if enrolled > 0:
            flash(f"Cannot delete — {enrolled} student(s) still enrolled.", "danger")
            return redirect(url_for("manage_courses"))

        subject_ids = [
            s.id for s in db.query(Subject).filter_by(course_id=course_id).all()
        ]
        semester_ids = [
            s.id for s in db.query(Semester).filter_by(course_id=course_id).all()
        ]

        if subject_ids:
            db.query(KnecMark).filter(KnecMark.subject_id.in_(subject_ids)).delete(synchronize_session=False)
            db.query(ShortCourseMark).filter(ShortCourseMark.subject_id.in_(subject_ids)).delete(synchronize_session=False)
            db.query(PeriodicMark).filter(PeriodicMark.subject_id.in_(subject_ids)).delete(synchronize_session=False)

        if semester_ids:
            db.query(KnecMark).filter(KnecMark.semester_id.in_(semester_ids)).delete(synchronize_session=False)
            db.query(ShortCourseMark).filter(ShortCourseMark.semester_id.in_(semester_ids)).delete(synchronize_session=False)
            db.query(PeriodicMark).filter(PeriodicMark.semester_id.in_(semester_ids)).delete(synchronize_session=False)

        _delete_quizzes_for(db, course_id=course_id)

        db.query(Message).filter(Message.course_id == course_id).delete(synchronize_session=False)
        if subject_ids:
            db.query(Message).filter(Message.subject_id.in_(subject_ids)).delete(synchronize_session=False)

        db.query(Subject).filter(Subject.course_id == course_id).delete(synchronize_session=False)
        db.query(Semester).filter(Semester.course_id == course_id).delete(synchronize_session=False)
        db.query(Module).filter(Module.course_id == course_id).delete(synchronize_session=False)

        db.delete(course)
        db.commit()
        flash("Course deleted successfully.", "success")
    except Exception as e:
        db.rollback()
        import traceback
        traceback.print_exc()
        flash(f"Error deleting course: {str(e)}", "danger")
    finally:
        db.close()

    return redirect(url_for("manage_courses"))


@app.route("/admin/course/<int:course_id>/modules", methods=["GET", "POST"])
def admin_course_modules(course_id):
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        course = db.query(Course).filter_by(id=course_id).first()
        if not course:
            flash("Course not found.", "danger")
            return redirect(url_for("manage_courses"))

        if request.method == "POST":
            action = request.form.get("action")

            if action == "add":
                name = request.form.get("name", "").strip()
                if not name:
                    flash("Module name is required.", "danger")
                elif db.query(Module).filter_by(course_id=course_id, name=name).first():
                    flash(f"'{name}' already exists for this course.", "warning")
                else:
                    last = (
                        db.query(Module).filter_by(course_id=course_id)
                        .order_by(Module.order.desc()).first()
                    )
                    next_order = (last.order + 1) if last else 1
                    db.add(Module(course_id=course_id, name=name, order=next_order))
                    db.commit()
                    flash(f"✅ Module '{name}' added.", "success")

            elif action == "rename":
                module_id = request.form.get("module_id", type=int)
                new_name = request.form.get("new_name", "").strip()
                mod = db.query(Module).filter_by(id=module_id, course_id=course_id).first()
                if mod and new_name:
                    mod.name = new_name
                    db.commit()
                    flash("✅ Module renamed.", "success")

            elif action == "delete":
                module_id = request.form.get("module_id", type=int)
                mod = db.query(Module).filter_by(id=module_id, course_id=course_id).first()
                if mod:
                    db.query(Subject).filter_by(module_id=module_id).update({"module_id": None})
                    db.query(StudentProfile).filter_by(module_id=module_id).update({"module_id": None})
                    db.delete(mod)
                    db.commit()
                    flash("✅ Module deleted.", "success")

            return redirect(url_for("admin_course_modules", course_id=course_id))

        modules = (
            db.query(Module).filter_by(course_id=course_id)
            .order_by(Module.order.asc(), Module.name.asc()).all()
        )
        mod_stats = [{
            "module": m,
            "subjects": db.query(Subject).filter_by(module_id=m.id).count(),
            "students": db.query(StudentProfile).filter_by(module_id=m.id).count(),
        } for m in modules]

        return render_template("admin/course_modules.html", course=course, mod_stats=mod_stats)
    finally:
        db.close()


@app.route("/add_subject", methods=["GET", "POST"])
def add_subject():
    blocked = require_admin_or_hod()
    if blocked:
        return blocked

    db = SessionLocal()
    try:
        if is_hod():
            hod_course_ids = [
                c.id for c in db.query(Course).filter_by(department_id=session["department_id"]).all()
            ]
        else:
            hod_course_ids = None

        if request.method == "POST":
            subject_name = request.form["subject_name"].strip()
            course_id = int(request.form["course_id"])
            semester_id = request.form.get("semester_id") or None
            module_id = request.form.get("module_id", type=int)

            if semester_id:
                semester_id = int(semester_id)

            if is_hod() and course_id not in hod_course_ids:
                flash("That course is not in your department.", "danger")
                return redirect(url_for("add_subject"))

            q = db.query(Subject).filter_by(name=subject_name, course_id=course_id)
            if semester_id:
                q = q.filter_by(semester_id=semester_id)
            else:
                q = q.filter(Subject.semester_id.is_(None))
            if q.first():
                flash("Subject already exists for this course/semester!", "danger")
                return redirect(url_for("add_subject"))

            if semester_id:
                sem = db.query(Semester).filter_by(id=semester_id).first()
                if not sem or sem.course_id != course_id:
                    flash("Selected semester does not belong to that course.", "danger")
                    return redirect(url_for("add_subject"))

            module_obj = None
            if module_id:
                module_obj = db.query(Module).filter_by(id=module_id, course_id=course_id).first()

            db.add(Subject(
                name=subject_name, course_id=course_id,
                semester_id=semester_id,
                module=module_obj.name if module_obj else None,
                module_id=module_obj.id if module_obj else None,
            ))
            db.commit()
            flash("Subject added successfully!", "success")
            return redirect(url_for("manage_courses"))

        if is_hod():
            courses = db.query(Course).filter_by(department_id=session["department_id"]).order_by(Course.name.asc()).all()
            semesters = (
                semester_query(db).filter(Semester.course_id.in_(hod_course_ids))
                .order_by(Semester.created_at.desc()).all()
                if hod_course_ids else []
            )
        else:
            courses = db.query(Course).order_by(Course.name.asc()).all()
            semesters = semester_query(db).order_by(Semester.created_at.desc()).all()

        modules = db.query(Module).order_by(Module.course_id.asc(), Module.order.asc()).all()

        return render_template(
            "admin/add_subject.html",
            courses=courses, semesters=semesters, modules=modules,
        )
    finally:
        db.close()


@app.route("/edit_subject/<int:subject_id>", methods=["GET", "POST"])
def edit_subject(subject_id):
    blocked = require_admin_or_hod()
    if blocked:
        return blocked

    db = SessionLocal()
    try:
        subject = subject_query(db).filter(Subject.id == subject_id).first()
        if not subject:
            return "Subject not found", 404

        if is_hod():
            course = db.query(Course).filter_by(
                id=subject.course_id, department_id=session["department_id"]
            ).first()
            if not course:
                flash("That subject is not in your department.", "danger")
                return redirect(url_for("manage_courses"))

        if request.method == "POST":
            subject.name = request.form["subject_name"].strip()

            semester_id = request.form.get("semester_id") or None
            if semester_id:
                semester_id = int(semester_id)
                sem = db.query(Semester).filter_by(id=semester_id).first()
                if not sem or sem.course_id != subject.course_id:
                    flash("Selected semester does not belong to that course.", "danger")
                    return redirect(url_for("edit_subject", subject_id=subject_id))
            subject.semester_id = semester_id

            module_id = request.form.get("module_id", type=int)
            if module_id:
                module_obj = db.query(Module).filter_by(id=module_id, course_id=subject.course_id).first()
                if module_obj:
                    subject.module_id = module_obj.id
                    subject.module = module_obj.name
            else:
                subject.module_id = None
                subject.module = None

            db.commit()
            flash("Subject updated.", "success")
            return redirect(url_for("manage_courses"))

        semesters = (
            semester_query(db).filter_by(course_id=subject.course_id)
            .order_by(Semester.created_at.desc()).all()
        )
        modules = (
            db.query(Module).filter_by(course_id=subject.course_id)
            .order_by(Module.order.asc(), Module.name.asc()).all()
        )
        return render_template(
            "admin/edit_subject.html",
            subject=subject, semesters=semesters, modules=modules,
        )
    finally:
        db.close()


@app.route("/delete_subject/<int:subject_id>", methods=["POST"])
def delete_subject(subject_id):
    blocked = require_admin_or_hod()
    if blocked:
        return blocked

    db = SessionLocal()
    try:
        subject = db.query(Subject).filter_by(id=subject_id).first()
        if not subject:
            return "Subject not found", 404

        if is_hod():
            course = db.query(Course).filter_by(
                id=subject.course_id, department_id=session["department_id"]
            ).first()
            if not course:
                flash("That subject is not in your department.", "danger")
                return redirect(url_for("manage_courses"))

        _delete_quizzes_for(db, subject_id=subject_id)
        db.query(Message).filter(Message.subject_id == subject_id).delete(synchronize_session=False)
        db.query(KnecMark).filter(KnecMark.subject_id == subject_id).delete(synchronize_session=False)
        db.query(ShortCourseMark).filter(ShortCourseMark.subject_id == subject_id).delete(synchronize_session=False)
        db.query(PeriodicMark).filter(PeriodicMark.subject_id == subject_id).delete(synchronize_session=False)
        db.delete(subject)
        db.commit()
        flash("Subject deleted successfully.", "success")
    except Exception as e:
        db.rollback()
        flash(f"Error deleting subject: {str(e)}", "danger")
    finally:
        db.close()

    return redirect(url_for("manage_courses"))


@app.route("/admin/semesters", methods=["GET"])
def manage_semesters():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        courses = db.query(Course).order_by(Course.name.asc()).all()
        semesters = semester_query(db).order_by(Semester.created_at.desc()).all()
        selected_course_id = request.args.get("course_id", type=int)

        return render_template(
            "admin/manage_semesters.html",
            courses=courses,
            semesters=semesters,
            selected_course_id=selected_course_id
        )
    finally:
        db.close()


@app.route("/admin/semesters/add", methods=["POST"])
def add_semester():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    name = request.form.get("name", "").strip()
    course_id = request.form.get("course_id")

    if not name or not course_id:
        flash("Semester name and course are required.", "danger")
        return redirect(url_for("manage_semesters"))

    db = SessionLocal()
    try:
        if db.query(Semester).filter_by(name=name, course_id=int(course_id)).first():
            flash(f"Semester '{name}' already exists for that course.", "warning")
            return redirect(url_for("manage_semesters"))

        db.add(Semester(name=name, course_id=int(course_id)))
        db.commit()
        flash("Semester added.", "success")
    except Exception as e:
        db.rollback()
        flash(f"Error: {e}", "danger")
    finally:
        db.close()

    return redirect(url_for("manage_semesters"))


@app.route("/admin/semesters/set_current/<int:semester_id>", methods=["POST"])
def set_current_semester(semester_id):
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        sem = db.query(Semester).filter_by(id=semester_id).first()
        if not sem:
            flash("Semester not found.", "danger")
            return redirect(url_for("manage_semesters"))

        db.query(Semester).filter(
            Semester.course_id == sem.course_id,
            Semester.id != sem.id,
        ).update({"is_current": False}, synchronize_session=False)

        sem.is_current = True
        db.commit()
        flash(f"'{sem.name}' is now current for its course.", "success")
    except Exception as e:
        db.rollback()
        flash(f"Error: {e}", "danger")
    finally:
        db.close()

    return redirect(url_for("manage_semesters"))


@app.route("/admin/semesters/delete/<int:semester_id>", methods=["POST"])
def delete_semester(semester_id):
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        sem = db.query(Semester).filter_by(id=semester_id).first()
        if not sem:
            flash("Semester not found.", "danger")
            return redirect(url_for("manage_semesters"))

        if db.query(Subject).filter_by(semester_id=semester_id).count():
            flash("Cannot delete — subjects are still in this semester.", "danger")
            return redirect(url_for("manage_semesters"))

        db.delete(sem)
        db.commit()
        flash("Semester deleted.", "success")
    except Exception as e:
        db.rollback()
        flash(f"Error: {e}", "danger")
    finally:
        db.close()

    return redirect(url_for("manage_semesters"))


@app.route("/upload_exam", methods=["GET", "POST"])
def upload_exam():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        if request.method == "POST":
            title = request.form["title"].strip()
            course_id = int(request.form["course"])
            subject_id = int(request.form["subject"])
            duration = int(request.form.get("duration", 30))
            file = request.files.get("quiz_file")

            if not file or not allowed(file.filename, ALLOWED_EXAM_EXTENSIONS):
                flash("❌ Upload a valid .docx file.", "danger")
                return redirect(url_for("upload_exam"))

            filename = secure_filename(file.filename)
            file_path = os.path.join(EXAMS_UPLOAD_FOLDER, filename)
            file.save(file_path)

            file.close()
            del file

            try:
                questions = parse_docx_questions(
                    file_path, image_output_dir=QUESTION_IMAGES_FOLDER
                )
            except Exception as parse_err:
                flash(f"❌ Could not parse file: {parse_err}", "danger")
                return redirect(url_for("upload_exam"))

            if not questions:
                flash("❌ No valid questions found.", "danger")
                return redirect(url_for("upload_exam"))

            quiz = Quiz(
                title=title, course_id=course_id, subject_id=subject_id,
                duration=duration, status="active",
            )
            db.add(quiz)
            db.commit()
            db.refresh(quiz)
            quiz_id = quiz.id

            for q in questions:
                db.add(Question(
                    quiz_id=quiz_id,
                    question_text=q.get("question", ""),
                    option_a=q.get("a", ""), option_b=q.get("b", ""),
                    option_c=q.get("c", ""), option_d=q.get("d", ""),
                    correct_option=q.get("answer", "").lower(),
                    marks=q.get("marks", 1),
                    extra_content=q.get("extra_content"),
                    image=q.get("image"),
                ))
            db.commit()
            flash(f"✅ Uploaded quiz with {len(questions)} question(s).", "success")

            del questions
            return redirect(url_for("manage_quizzes"))

        courses = db.query(Course).all()
        subjects = subject_query(db).all()
        return render_template(
            "admin/upload_exams.html", courses=courses, subjects=subjects
        )
    except Exception as e:
        db.rollback()
        import traceback
        traceback.print_exc()
        flash(f"❌ Upload failed: {e}", "danger")
        return redirect(url_for("upload_exam"))
    finally:
        db.close()

@app.route("/admin/manage_quizzes", methods=["GET", "POST"])
def manage_quizzes():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        if request.method == "POST":
            quiz_id = int(request.form.get("quiz_id"))
            action = request.form.get("action")

            quiz = db.query(Quiz).filter_by(id=quiz_id).first()
            if not quiz:
                flash("Quiz not found.", "danger")
                return redirect(url_for("manage_quizzes"))

            if action == "activate":
                quiz.status = "active"
            elif action == "deactivate":
                quiz.status = "inactive"

            db.commit()
            flash(f"Quiz '{quiz.title}' has been {quiz.status}.", "success")
            return redirect(url_for("manage_quizzes"))

        quizzes = quiz_query(db).all()
        return render_template("admin/manage_quizzes.html", quizzes=quizzes)
    finally:
        db.close()

@app.route("/admin/delete_quiz/<int:quiz_id>", methods=["POST"])
def delete_quiz(quiz_id):
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        quiz = db.query(Quiz).filter_by(id=quiz_id).first()
        if quiz:
            _delete_quiz_tree(db, quiz)
            db.commit()
            flash("Quiz deleted successfully.", "success")
        else:
            flash("Quiz not found.", "danger")
    except Exception as e:
        db.rollback()
        flash(f"Error deleting quiz: {str(e)}", "danger")
    finally:
        db.close()

    return redirect(url_for("view_results"))

@app.route("/admin/review_quiz/<int:quiz_id>", methods=["GET", "POST"])
def review_uploaded_quiz(quiz_id):
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        quiz = quiz_query(db).filter_by(id=quiz_id).first()
        if not quiz:
            flash("Quiz not found.", "danger")
            return redirect(url_for("upload_exam"))

        if request.method == "POST":
            action = request.form.get("action")
            if action == "delete":
                _delete_quiz_tree(db, quiz)
                db.commit()
                flash("❌ Quiz deleted.", "warning")
                return redirect(url_for("upload_exam"))
            elif action == "confirm":
                flash("✅ Quiz confirmed and saved.", "success")
                return redirect(url_for("upload_exam"))

        questions = db.query(Question).filter_by(quiz_id=quiz.id).all()
        return render_template(
            "admin/review_uploaded_quiz.html", quiz=quiz, questions=questions
        )
    finally:
        db.close()


# @app.route("/manage_students", methods=["GET", "POST"])
# def manage_students():
#     if session.get("role") != "admin":
#         return redirect(url_for("login"))

#     db = SessionLocal()
#     try:
#         if request.method == "POST":
#             student_id = request.form.get("student_id")
#             action = request.form.get("action")
#             student = db.query(StudentProfile).filter_by(id=student_id).first()

#             if action == "toggle_block" and student:
#                 student.blocked = not student.blocked
#                 db.commit()
#             elif action == "delete" and student:
#                 try:
#                     _delete_student_tree(db, student)
#                     db.commit()
#                     flash("✅ Student deleted.", "success")
#                 except Exception as e:
#                     db.rollback()
#                     flash(f"Error deleting student: {str(e)}", "danger")

#             return redirect(url_for("manage_students"))

#         search_query_str = request.args.get("search", "").strip()
#         if search_query_str:
#             students = (
#                 student_query(db)
#                 .filter(or_(
#                     StudentProfile.full_name.ilike(f"%{search_query_str}%"),
#                     StudentProfile.exam_type.ilike(f"%{search_query_str}%"),
#                     StudentProfile.admission_number.ilike(f"%{search_query_str}%"),
#                     StudentProfile.phone_number.ilike(f"%{search_query_str}%"),
#                 ))
#                 .all()
#             )
#         else:
#             students = student_query(db).all()

#         return render_template(
#             "admin/manage_students.html",
#             students=students, search_query=search_query_str,
#         )
#     finally:
#         db.close()



@app.route("/admin/students", methods=["GET", "POST"])
def admin_all_students():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        dept_id = request.args.get("dept_id", type=int)
        course_id = request.args.get("course_id", type=int)
        search = request.args.get("search", "").strip()

        departments = db.query(Department).order_by(Department.name.asc()).all()
        courses = []
        students = []

        if dept_id:
            courses = db.query(Course).filter_by(department_id=dept_id).order_by(Course.name.asc()).all()
            q = student_query(db)
            if course_id:
                q = q.filter(StudentProfile.course_id == course_id)
            else:
                course_ids = [c.id for c in courses]
                q = q.filter(StudentProfile.course_id.in_(course_ids))
            if search:
                like = f"%{search}%"
                q = q.filter(or_(
                    StudentProfile.full_name.ilike(like),
                    StudentProfile.admission_number.ilike(like),
                    StudentProfile.phone_number.ilike(like),
                ))
            students = q.order_by(StudentProfile.full_name.asc()).all()

        if request.method == "POST":
            student_id = request.form.get("student_id")
            action = request.form.get("action")
            student = db.query(StudentProfile).filter_by(id=student_id).first()

            if action == "toggle_block" and student:
                student.blocked = not student.blocked
                db.commit()
                flash("Student status updated.", "success")
            elif action == "delete" and student:
                _delete_student_tree(db, student)
                db.commit()
                flash("Student deleted.", "success")

            return redirect(url_for("admin_all_students", dept_id=dept_id, course_id=course_id, search=search))

        return render_template(
            "admin/all_students.html",
            departments=departments, courses=courses,
            students=students, dept_id=dept_id, course_id=course_id, search=search
        )
    finally:
        db.close()

@app.route("/admin/students/add", methods=["GET", "POST"])
def admin_add_student():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        departments = db.query(Department).order_by(Department.name.asc()).all()
        courses = db.query(Course).order_by(Course.name.asc()).all()

        if request.method == "POST":
            data, errors = _validate_new_student_form(db, request.form)

            if errors:
                for e in errors:
                    flash(f"❌ {e}", "danger")

                # ── C2: repopulate dependent dropdowns for the re-render ──
                error_course_id = request.form.get("course_id", type=int)
                error_semesters = []
                error_modules = []
                if error_course_id:
                    error_semesters = (
                        db.query(Semester)
                        .filter_by(course_id=error_course_id)
                        .order_by(Semester.created_at.asc())
                        .all()
                    )
                    error_modules = (
                        db.query(Module)
                        .filter_by(course_id=error_course_id)
                        .order_by(Module.order.asc(), Module.name.asc())
                        .all()
                    )

                return render_template(
                    "admin/add_student.html",
                    departments=departments,
                    courses=courses,
                    semesters=error_semesters,
                    modules=error_modules,
                    form=request.form,
                    errors=errors,
                    error_fields=_error_fields_from(errors),
                    username=session.get("username"),
                )

            ok, msg, _ = _create_single_student(
                db,
                full_name=data["full_name"],
                admission=data["admission"],
                course=data["course"],
                exam_type=data["exam_type"],
                semester_id=data["semester_id"],
                module_obj=data["module_obj"],
                phone=data["phone"],
                username=data["username"],
                password=data["password"],
            )

            if not ok:
                db.rollback()
                flash(f"❌ {msg}", "danger")
                return render_template(
                    "admin/add_student.html",
                    departments=departments,
                    courses=courses,
                    semesters=[],
                    modules=[],
                    form=request.form,
                    errors=[msg],
                    error_fields=set(),
                    username=session.get("username"),
                )

            db.commit()
            flash(f"✅ {msg}", "success")
            return redirect(url_for("admin_all_students"))

        # ── GET branch ──
        return render_template(
            "admin/add_student.html",
            departments=departments,
            courses=courses,
            semesters=[],
            modules=[],
            form={},
            errors=[],
            error_fields=set(),
            username=session.get("username"),
        )

    except IntegrityError as e:
        db.rollback()
        traceback.print_exc()
        flash(f"❌ {_translate_db_error(e)}", "danger")
        return redirect(url_for("admin_add_student"))

    except Exception:
        db.rollback()
        traceback.print_exc()
        flash("❌ An unexpected error occurred. Please try again.", "danger")
        return redirect(url_for("admin_add_student"))

    finally:
        db.close()



@app.route("/admin/student/<int:student_id>/edit", methods=["GET", "POST"])
@app.route("/hod/student/<int:student_id>/edit", methods=["GET", "POST"])
def edit_student(student_id):
    role = session.get("role")
    if role not in ("admin", "hod"):
        return redirect(url_for("login"))

    db = SessionLocal()
    back_url = url_for("login")
    try:
        student = db.query(StudentProfile).filter_by(id=student_id).first()
        if not student:
            flash("❌ Student not found.", "danger")
            return redirect(url_for("login"))

        if role == "hod":
            course = db.query(Course).filter_by(id=student.course_id).first()
            if not course or course.department_id != session["department_id"]:
                flash("❌ That student is not in your department.", "danger")
                return redirect(url_for("hod_dashboard"))

            available_courses = db.query(Course).filter_by(department_id=session["department_id"]).all()
            available_semesters = db.query(Semester).filter(Semester.course_id.in_([c.id for c in available_courses])).all()
            available_modules = db.query(Module).filter(Module.course_id.in_([c.id for c in available_courses])).all()
            back_url = url_for("hod_students_list")
        else:
            available_courses = db.query(Course).all()
            available_semesters = db.query(Semester).all()
            available_modules = db.query(Module).all()
            back_url = url_for("manage_students")

        if request.method == "POST":
            student.full_name = request.form.get("full_name", "").strip()
            student.admission_number = request.form.get("admission_number", "").strip()
            student.phone_number = request.form.get("phone_number", "").strip() or None

            new_course_id = request.form.get("course_id", type=int)
            new_semester_id = request.form.get("semester_id", type=int)
            new_module_id = request.form.get("module_id", type=int)

            if new_course_id:
                student.course_id = new_course_id

            if new_semester_id:
                sem = db.query(Semester).filter_by(id=new_semester_id).first()
                if sem and sem.course_id == student.course_id:
                    student.current_semester_id = new_semester_id
                else:
                    student.current_semester_id = None
            else:
                student.current_semester_id = None

            if new_module_id:
                mod = db.query(Module).filter_by(id=new_module_id).first()
                if mod and mod.course_id == student.course_id:
                    student.module_id = new_module_id
                    student.module = mod.name
                else:
                    student.module_id = None
                    student.module = None
            else:
                student.module_id = None
                student.module = None

            db.commit()
            flash(f"✅ {student.full_name} updated successfully.", "success")
            return redirect(back_url)

        return render_template(
            "admin/edit_student.html",
            student=student,
            courses=available_courses,
            semesters=available_semesters,
            modules=available_modules,
            back_url=back_url
        )
    except Exception as e:
        db.rollback()
        flash(f"❌ Error updating student: {e}", "danger")
        return redirect(back_url)
    finally:
        db.close()


@app.route("/admin/assign-modules", methods=["GET", "POST"])
def admin_assign_modules():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        courses = db.query(Course).order_by(Course.name.asc()).all()

        if request.method == "POST":
            course_id = request.form.get("course_id", type=int)
            if not course_id:
                flash("Pick a course first.", "danger")
                return redirect(url_for("admin_assign_modules"))

            course = db.query(Course).filter_by(id=course_id).first()
            if not course:
                flash("Course not found.", "danger")
                return redirect(url_for("admin_assign_modules"))

            students = (
                db.query(StudentProfile).filter_by(course_id=course_id)
                .order_by(StudentProfile.full_name.asc()).all()
            )

            updated = 0
            for s in students:
                field = f"module_id_{s.id}"
                raw = request.form.get(field, "").strip()

                if raw == "":
                    if s.module_id is not None:
                        s.module_id = None
                        s.module = None
                        updated += 1
                    continue

                module_id = int(raw)
                mod = db.query(Module).filter_by(id=module_id, course_id=course_id).first()
                if not mod:
                    continue

                if s.module_id != mod.id:
                    s.module_id = mod.id
                    s.module = mod.name
                    updated += 1

            db.commit()
            flash(f"✅ Updated {updated} student(s).", "success")
            return redirect(url_for("admin_assign_modules", course_id=course_id))

        course_id = request.args.get("course_id", type=int)
        course = None
        students = []
        modules = []

        if course_id:
            course = db.query(Course).filter_by(id=course_id).first()
            if course:
                students = (
                    db.query(StudentProfile).filter_by(course_id=course_id)
                    .order_by(StudentProfile.full_name.asc()).all()
                )
                modules = (
                    db.query(Module).filter_by(course_id=course_id)
                    .order_by(Module.order.asc(), Module.name.asc()).all()
                )

        return render_template(
            "admin/assign_modules.html",
            courses=courses, course=course,
            students=students, modules=modules,
        )
    finally:
        db.close()


@app.route("/admin/students/upload/template")
def admin_students_template():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        courses = db.query(Course).order_by(Course.name.asc()).all()

        sample = pd.DataFrame([{
            "full_name": "silas sakwa",
            "admission_number": "STC/2025/008",
            "course_name": courses[0].name if courses else "CERTIFICATE IN ICT",
            "department_name": courses[0].department.name if courses and courses[0].department else "",
            "exam_type": "knec",
            "module": "MODULE 1",
            "semester": "OCT-DEC 2025",
            "phone_number": "0745361425",
            "username": "silas",
            "password": "1234",
        }])

        ref_rows = []
        for c in courses:
            sems = db.query(Semester).filter_by(course_id=c.id).order_by(Semester.created_at.asc()).all()
            dept_name = c.department.name if c.department else ""
            if sems:
                for s in sems:
                    ref_rows.append({
                        "department_name": dept_name,
                        "course_name": c.name,
                        "level": c.level,
                        "semester (term)": s.name,
                    })
            else:
                ref_rows.append({
                    "department_name": dept_name,
                    "course_name": c.name,
                    "level": c.level,
                    "semester (term)": "(no semesters yet)",
                })

        reference = pd.DataFrame(
            ref_rows or [{"department_name": "", "course_name": "(no courses yet)", "level": "", "semester (term)": ""}]
        )

        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            sample.to_excel(writer, index=False, sheet_name="Students")
            reference.to_excel(writer, index=False, sheet_name="Valid Courses")
            for name, df in (("Students", sample), ("Valid Courses", reference)):
                ws = writer.sheets[name]
                for idx, col in enumerate(df.columns, start=1):
                    max_len = max([len(str(col))] + [len(str(v)) for v in df[col].tolist()]) if len(df) else len(str(col))
                    ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(max(max_len + 3, 12), 45)
        buf.seek(0)

        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True, download_name="student_upload_template.xlsx",
        )
    finally:
        db.close()


@app.route("/admin/students/upload", methods=["GET", "POST"])
def admin_students_upload():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        departments = db.query(Department).order_by(Department.name.asc()).all()

        if request.method == "POST":
            dept_id = request.form.get("department_id", type=int)
            if not dept_id:
                flash("❌ Please choose a department first.", "danger")
                return redirect(url_for("admin_students_upload"))

            department = db.query(Department).filter_by(id=dept_id).first()
            if not department:
                flash("Department not found.", "danger")
                return redirect(url_for("admin_students_upload"))

            file = request.files.get("student_file")
            if not file or not file.filename.lower().endswith((".xlsx", ".xls")):
                flash("❌ Please upload a valid .xlsx or .xls file.", "danger")
                return redirect(url_for("admin_students_upload"))

            try:
                df = pd.read_excel(file, sheet_name="Students", dtype=str).fillna("")
            except Exception as e:
                flash(f"❌ Could not read the Excel file: {e}", "danger")
                return redirect(url_for("admin_students_upload"))

            df.columns = [str(c).strip().lower().replace(" ", "_") for c in df.columns]

            required_cols = {"full_name", "admission_number", "course_name", "exam_type"}
            missing = required_cols - set(df.columns)
            if missing:
                flash(f"❌ Missing required columns: {', '.join(sorted(missing))}", "danger")
                return redirect(url_for("admin_students_upload"))

            allowed_courses = db.query(Course).filter_by(department_id=dept_id).all()
            course_by_name = {c.name.strip().lower(): c for c in allowed_courses}

            existing_adm = {a[0] for a in db.query(StudentProfile.admission_number).all() if a[0]}
            seen_in_file = set()
            valid_rows, error_rows = [], []

            for idx, row in df.iterrows():
                row_no = idx + 2
                full_name = str(row.get("full_name", "")).strip()
                admission = str(row.get("admission_number", "")).strip()
                course_name = str(row.get("course_name", "")).strip()
                exam_type = str(row.get("exam_type", "")).strip().lower()
                module_name = str(row.get("module", "")).strip()
                semester_name = str(row.get("semester", "")).strip()
                phone = str(row.get("phone_number", "")).strip()
                username = str(row.get("username", "")).strip()
                password = str(row.get("password", "")).strip()

                errors = []
                semester_id = None

                if not full_name:
                    errors.append("full_name is empty")
                if not admission:
                    errors.append("admission_number is empty")
                elif admission in existing_adm:
                    errors.append(f"admission_number '{admission}' already exists")
                elif admission in seen_in_file:
                    errors.append(f"admission_number '{admission}' duplicated inside the file")
                else:
                    seen_in_file.add(admission)

                matched_course = course_by_name.get(course_name.lower())
                if not matched_course:
                    errors.append(f"course '{course_name}' not found in department '{department.name}'")
                else:
                    if get_course_assessment_type(matched_course) in ("short", "periodic"):
                        if exam_type != "internal":
                            exam_type = "internal"

                    if semester_name:
                        sem = db.query(Semester).filter(
                            Semester.course_id == matched_course.id,
                            Semester.name.ilike(semester_name),
                        ).first()
                        if not sem:
                            errors.append(f"semester '{semester_name}' not found for course '{course_name}'")
                        else:
                            semester_id = sem.id

                if exam_type not in ("knec", "internal"):
                    errors.append("exam_type must be 'knec' or 'internal'")

                if errors:
                    error_rows.append({
                        "row": row_no, "full_name": full_name,
                        "admission_number": admission, "errors": errors,
                    })
                else:
                    valid_rows.append({
                        "row": row_no,
                        "full_name": full_name,
                        "admission_number": admission,
                        "course_name": course_name,
                        "course_id": matched_course.id,
                        "exam_type": exam_type,
                        "module": module_name or None,
                        "semester_name": semester_name or None,
                        "semester_id": semester_id,
                        "phone_number": phone or None,
                        "username": username or admission.replace("/", "").replace(" ", "").lower(),
                        "password": password or admission,
                    })

            session["admin_upload_valid"] = valid_rows
            session["admin_upload_dept"] = dept_id

            return render_template(
                "admin/upload_students_preview.html",
                department=department,
                valid_rows=valid_rows, error_rows=error_rows, total_rows=len(df),
            )

        return render_template("admin/upload_students.html", departments=departments)
    finally:
        db.close()


@app.route("/admin/students/upload/confirm", methods=["POST"])
def admin_students_upload_confirm():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    rows = session.pop("admin_upload_valid", None)
    dept_id = session.pop("admin_upload_dept", None)
    if not rows:
        flash("No validated rows to import. Please upload again.", "warning")
        return redirect(url_for("admin_students_upload"))

    db = SessionLocal()
    imported, skipped = 0, 0
    try:
        course_ids = list({r["course_id"] for r in rows})
        courses = {
            c.id: c for c in db.query(Course).filter(
                Course.id.in_(course_ids),
                Course.department_id == dept_id,
            ).all()
        }

        admission_numbers = [r["admission_number"] for r in rows]
        existing_admissions = {
            a[0] for a in db.query(StudentProfile.admission_number)
                           .filter(StudentProfile.admission_number.in_(admission_numbers))
                           .all()
            if a[0]
        }

        base_usernames = [(r["username"] or r["admission_number"]).strip()[:45] for r in rows]
        existing_usernames = {
            u[0] for u in db.query(User.username)
                           .filter(User.username.in_(base_usernames))
                           .all()
            if u[0]
        }

        current_semesters = {
            s.course_id: s.id
            for s in db.query(Semester).filter(
                Semester.course_id.in_(course_ids),
                Semester.is_current == True,
            ).all()
        }

        modules_lookup = {}
        for m in db.query(Module).filter(Module.course_id.in_(course_ids)).all():
            modules_lookup[(m.course_id, m.name)] = m

        new_users, pending = [], []
        username_taken = set(existing_usernames)

        for r in rows:
            course = courses.get(r["course_id"])
            if not course:
                skipped += 1
                continue

            adm = r["admission_number"]
            if adm in existing_admissions:
                skipped += 1
                continue

            base_username = (r["username"] or adm).strip()[:45] or adm
            username = base_username
            counter = 1
            while username in username_taken:
                username = f"{base_username}{counter}"
                counter += 1
            username_taken.add(username)

            user = User(username=username, password=r["password"] or adm)

            sem_id = r.get("semester_id") or current_semesters.get(course.id)

            module_id = None
            module_label = r.get("module")
            if module_label:
                mod = modules_lookup.get((course.id, module_label))
                if mod:
                    module_id = mod.id
                    module_label = mod.name

            new_users.append(user)
            pending.append({
                "course_id": course.id,
                "full_name": r["full_name"],
                "admission_number": adm,
                "exam_type": r["exam_type"],
                "phone_number": r["phone_number"],
                "current_semester_id": sem_id,
                "module": module_label,
                "module_id": module_id,
            })
            existing_admissions.add(adm)

        if new_users:
            db.add_all(new_users)
            db.flush()
            new_students = [StudentProfile(user_id=user.id, **data) for user, data in zip(new_users, pending)]
            db.add_all(new_students)
            db.commit()
            imported = len(new_students)
        else:
            db.commit()

        msg = f"✅ Imported {imported} student(s)."
        if skipped:
            msg += f" Skipped {skipped} (duplicates or course mismatch)."
        flash(msg, "success" if imported else "warning")

    except IntegrityError as e:
        db.rollback()
        import traceback; traceback.print_exc()
        flash(f"❌ Import failed: {_translate_db_error(e)}", "danger")
    except Exception as e:
        db.rollback()
        import traceback
        traceback.print_exc()
        flash(f"❌ Import failed: {e}", "danger")
    finally:
        db.close()

    return redirect(url_for("manage_students"))


@app.route("/show_credentials")
def show_credentials():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        users = db.query(User).options(joinedload(User.profile)).all()
        return render_template("admin/show_credentials.html", users=users)
    finally:
        db.close()


@app.route("/delete_user/<int:user_id>", methods=["POST"])
def delete_user(user_id):
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        user = db.query(User).filter(User.id == user_id).first()
        if user:
            profile = user.profile
            if profile:
                _delete_student_tree(db, profile)
            else:
                db.delete(user)
            db.commit()
            flash("✅ User deleted.", "success")
        else:
            flash("User not found.", "danger")
    except Exception as e:
        db.rollback()
        flash(f"Error deleting user: {str(e)}", "danger")
    finally:
        db.close()

    return redirect(url_for("show_credentials"))


@app.route("/admin/view_results", methods=["GET", "POST"])
def view_results():
    if "user_id" not in session or session.get("role") != "admin":
        flash("Admin access required", "danger")
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        courses = db.query(Course).all()
        subjects = subject_query(db).all()

        selected_course = request.form.get("course")
        selected_subject = request.form.get("subject")
        export = request.form.get("export")

        query = (
            db.query(Result)
            .options(
                joinedload(Result.student),
                joinedload(Result.quiz).joinedload(Quiz.course),
                joinedload(Result.quiz).joinedload(Quiz.subject),
            )
            .join(Result.quiz).join(Quiz.course).join(Quiz.subject).join(Result.student)
        )

        if selected_course:
            query = query.filter(Quiz.course_id == int(selected_course))
        if selected_subject:
            query = query.filter(Quiz.subject_id == int(selected_subject))

        results = query.all()

        if export == "true":
            data = []
            for r in results:
                data.append({
                    "Student Username": r.student.user.username if r.student and r.student.user else "N/A",
                    "Full Name": r.student.full_name if r.student else "N/A",
                    "Course": r.quiz.course.name if r.quiz.course else "N/A",
                    "Subject": r.quiz.subject.name if r.quiz.subject else "N/A",
                    "Quiz Title": r.quiz.title,
                    "Score": r.score,
                    "Total Marks": r.total_marks,
                    "Percentage": r.percentage,
                    "Taken On": r.taken_on.strftime("%Y-%m-%d %H:%M:%S"),
                })

            df = pd.DataFrame(data)
            excel_path = os.path.join(EXAMS_UPLOAD_FOLDER, "quiz_results.xlsx")
            df.to_excel(excel_path, index=False)

            return send_file(
                excel_path,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                as_attachment=True, download_name="quiz_results.xlsx",
            )

        return render_template(
            "admin/view_results.html",
            results=results, courses=courses, subjects=subjects,
            selected_course=selected_course, selected_subject=selected_subject,
        )
    finally:
        db.close()


@app.route("/admin/messages", methods=["GET", "POST"])
def admin_messages():
    if session.get("role") != "admin":
        flash("Please log in as admin.", "error")
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        if request.method == "POST":
            content = request.form["content"]
            target_type = request.form["target_type"]
            course_id = request.form.get("course_id") or None
            subject_id = request.form.get("subject_id") or None

            db.add(Message(
                content=content, target_type=target_type,
                course_id=course_id if target_type == "course" else None,
                subject_id=subject_id if target_type == "subject" else None,
            ))
            db.commit()
            flash("Message created successfully!", "success")
            return redirect(url_for("admin_messages"))

        courses = db.query(Course).all()
        subjects = subject_query(db).all()
        messages = db.query(Message).order_by(Message.created_at.desc()).all()

        return render_template(
            "admin/messages.html",
            courses=courses, subjects=subjects, messages=messages,
        )
    finally:
        db.close()


@app.route("/admin/messages/delete/<int:message_id>", methods=["POST"])
def delete_admin_message(message_id):
    if session.get("role") != "admin":
        flash("Admin access required.", "danger")
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        message = db.query(Message).filter_by(id=message_id).first()
        if not message:
            flash("Message not found.", "danger")
        else:
            db.delete(message)
            db.commit()
            flash("Message deleted successfully.", "success")
    finally:
        db.close()

    return redirect(url_for("admin_messages"))


@app.route("/admin/student_activity")
def student_activity():
    if session.get("role") != "admin":
        flash("Please log in as admin.", "error")
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        total_students = db.query(StudentProfile).count()
        total_doing_exam = (
            db.query(ActivityLog)
            .filter(ActivityLog.activity_type == "exam", ActivityLog.is_active == True)
            .count()
        )

        StudentActivity = namedtuple("StudentActivity", ["full_name", "course_name", "activity_type"])
        rows = (
            db.query(StudentProfile.full_name, Course.name.label("course_name"), ActivityLog.activity_type)
            .join(ActivityLog, ActivityLog.student_id == StudentProfile.id)
            .join(Course, StudentProfile.course_id == Course.id)
            .filter(ActivityLog.is_active == True, ActivityLog.activity_type == "exam")
            .all()
        )
        active_students = [StudentActivity(*row) for row in rows]

        return render_template(
            "admin/student_activity.html",
            total_students=total_students,
            total_doing_exam=total_doing_exam,
            active_students=active_students,
        )
    finally:
        db.close()


@app.route("/admin/marks", methods=["GET"])
def admin_marks_picker():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        departments = db.query(Department).order_by(Department.name.asc()).all()

        dept_cards = []
        for d in departments:
            course_ids = [c.id for c in db.query(Course).filter_by(department_id=d.id).all()]
            students_count = 0
            if course_ids:
                students_count = (
                    db.query(StudentProfile)
                    .filter(StudentProfile.course_id.in_(course_ids))
                    .count()
                )
            dept_cards.append({
                "department": d,
                "courses_count": len(course_ids),
                "students_count": students_count,
            })

        return render_template("admin/marks_step1_departments.html", dept_cards=dept_cards)
    finally:
        db.close()


@app.route("/admin/marks/department/<int:dept_id>", methods=["GET"])
def admin_marks_department(dept_id):
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        department = db.query(Department).filter_by(id=dept_id).first()
        if not department:
            flash("Department not found.", "danger")
            return redirect(url_for("admin_marks_picker"))

        courses = db.query(Course).filter_by(department_id=dept_id).order_by(Course.name.asc()).all()

        course_id = request.args.get("course_id", type=int)
        semester_id = request.args.get("semester_id", type=int)
        module_id = request.args.get("module_id", type=int)

        course = semester = None
        subjects, semesters, available_modules = [], [], []

        if course_id:
            course = db.query(Course).filter_by(id=course_id, department_id=dept_id).first()
        if course:
            semesters = (
                semester_query(db).filter_by(course_id=course.id)
                .order_by(Semester.created_at.desc()).all()
            )
            if semester_id:
                semester = next((s for s in semesters if s.id == semester_id), None)
            elif semesters:
                semester = next((s for s in semesters if s.is_current), semesters[0])

            if semester:
                all_sem_subjects = (
                    subject_query(db)
                    .filter_by(course_id=course.id, semester_id=semester.id)
                    .order_by(Subject.name.asc()).all()
                )
                used_module_ids = {s.module_id for s in all_sem_subjects if s.module_id}
                if used_module_ids:
                    available_modules = (
                        db.query(Module).filter(Module.id.in_(used_module_ids))
                        .order_by(Module.order.asc(), Module.name.asc()).all()
                    )
                subjects = (
                    [s for s in all_sem_subjects if s.module_id == module_id]
                    if module_id else all_sem_subjects
                )

        return render_template(
            "admin/marks_step2_course.html",
            department=department, courses=courses,
            course=course, semesters=semesters, semester=semester,
            subjects=subjects, available_modules=available_modules, module_id=module_id,
        )
    finally:
        db.close()


@app.route("/admin/marks/enter", methods=["GET", "POST"])
def admin_enter_marks():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        src = request.form if request.method == "POST" else request.args
        subject_id = src.get("subject_id", type=int)
        semester_id = src.get("semester_id", type=int)

        if not subject_id:
            flash("Pick a subject first.", "danger")
            return redirect(url_for("admin_marks_picker"))

        subject = subject_query(db).filter_by(id=subject_id).first()
        if not subject:
            flash("Subject not found.", "danger")
            return redirect(url_for("admin_marks_picker"))

        semester = None
        if semester_id:
            semester = db.query(Semester).filter_by(id=semester_id).first()
            if not semester:
                flash("Semester not found.", "danger")
                return redirect(url_for("admin_marks_picker"))
            if subject.course_id != semester.course_id:
                flash("Subject and semester belong to different courses.", "danger")
                return redirect(url_for("admin_marks_picker"))

        course = subject.course
        department = course.department if course else None
        atype = get_course_assessment_type(course)

        students_q = student_query(db).filter_by(course_id=subject.course_id)
        fresh_module_id = db.query(Subject.module_id).filter(Subject.id == subject_id).scalar()

        if fresh_module_id:
            students_q = students_q.filter(
                or_(
                    StudentProfile.module_id == fresh_module_id,
                    StudentProfile.module_id.is_(None),
                )
            )

        students = students_q.order_by(StudentProfile.full_name.asc()).all()

        if request.method == "POST":
            def to_float(v):
                v = (v or "").strip()
                if not v:
                    return None
                try:
                    return float(v)
                except ValueError:
                    return None

            def to_int(v):
                v = (v or "").strip()
                return int(v) if v.isdigit() else None

            saved = 0

            if atype == "short":
                max_s = course.max_score or 100
                for s in students:
                    sc = to_float(request.form.get(f"score_{s.id}"))
                    rm = (request.form.get(f"remarks_{s.id}") or "").strip() or None
                    if sc is not None and not (0 <= sc <= max_s):
                        flash(f"Score for {s.full_name} must be 0–{max_s}.", "warning")
                        continue
                    existing = db.query(ShortCourseMark).filter_by(
                        student_id=s.id, subject_id=subject_id, semester_id=semester_id
                    ).first()
                    if sc is None and not rm:
                        if existing:
                            db.delete(existing)
                        continue
                    if existing:
                        existing.score = sc
                        existing.remarks = rm
                        existing.entered_at = datetime.utcnow()
                    else:
                        db.add(ShortCourseMark(
                            student_id=s.id, subject_id=subject_id,
                            semester_id=semester_id, score=sc, remarks=rm,
                        ))
                    saved += 1

            elif atype == "kasneb":
                max_s = course.max_score or 100
                for s in students:
                    c1 = to_float(request.form.get(f"score_{s.id}__CAT 1"))
                    c2 = to_float(request.form.get(f"score_{s.id}__CAT 2"))
                    if c1 is not None and not (0 <= c1 <= max_s):
                        flash(f"CAT 1 for {s.full_name} must be 0–{max_s}.", "warning")
                        continue
                    if c2 is not None and not (0 <= c2 <= max_s):
                        flash(f"CAT 2 for {s.full_name} must be 0–{max_s}.", "warning")
                        continue
                    _save_kasneb_marks(db, s.id, subject_id, semester_id, c1, c2)
                    if c1 is not None or c2 is not None:
                        saved += 1

            elif atype == "periodic":
                max_s = course.max_score or 100
                labels = get_periodic_labels(course)
                for s in students:
                    for label in labels:
                        sc = to_float(request.form.get(f"score_{s.id}__{label}"))
                        rm = (request.form.get(f"remarks_{s.id}__{label}") or "").strip() or None
                        if sc is not None and not (0 <= sc <= max_s):
                            flash(f"'{label}' for {s.full_name} must be 0–{max_s}.", "warning")
                            continue
                        existing = db.query(PeriodicMark).filter_by(
                            student_id=s.id, subject_id=subject_id,
                            semester_id=semester_id, label=label,
                        ).first()
                        if sc is None and not rm:
                            if existing:
                                db.delete(existing)
                            continue
                        if existing:
                            existing.score = sc
                            existing.remarks = rm
                            existing.entered_at = datetime.utcnow()
                        else:
                            db.add(PeriodicMark(
                                student_id=s.id, subject_id=subject_id,
                                semester_id=semester_id, label=label,
                                score=sc, remarks=rm,
                            ))
                        saved += 1

            else:  # KNEC
                for s in students:
                    c1 = to_int(request.form.get(f"cat1_{s.id}"))
                    c2 = to_int(request.form.get(f"cat2_{s.id}"))
                    fn = to_int(request.form.get(f"final_{s.id}"))
                    existing = db.query(KnecMark).filter_by(
                        student_id=s.id, subject_id=subject_id, semester_id=semester_id
                    ).first()
                    if c1 is None and c2 is None and fn is None:
                        if existing:
                            db.delete(existing)
                        continue
                    if c1 is not None and not (0 <= c1 <= 15):
                        flash(f"CAT1 for {s.full_name} must be 0–15.", "warning")
                        continue
                    if c2 is not None and not (0 <= c2 <= 15):
                        flash(f"CAT2 for {s.full_name} must be 0–15.", "warning")
                        continue
                    if fn is not None and not (0 <= fn <= 70):
                        flash(f"Final for {s.full_name} must be 0–70.", "warning")
                        continue
                    if existing:
                        existing.cat1, existing.cat2, existing.final = c1, c2, fn
                        existing.entered_at = datetime.utcnow()
                    else:
                        db.add(KnecMark(
                            student_id=s.id, subject_id=subject_id,
                            semester_id=semester_id,
                            cat1=c1, cat2=c2, final=fn,
                        ))
                    saved += 1

            db.commit()
            flash(f"✅ Saved marks for {saved} record(s).", "success")

            if department:
                return redirect(url_for(
                    "admin_marks_department",
                    dept_id=department.id,
                    course_id=subject.course_id,
                    semester_id=semester_id,
                ))
            return redirect(url_for("admin_marks_picker"))

        if atype == "short":
            existing = {m.student_id: m for m in db.query(ShortCourseMark).filter_by(
                subject_id=subject_id, semester_id=semester_id).all()}
        elif atype in ("kasneb", "periodic"):
            existing = {}
            for m in db.query(PeriodicMark).filter_by(
                subject_id=subject_id, semester_id=semester_id).all():
                existing.setdefault(m.student_id, {})[m.label] = m
        else:
            existing = {m.student_id: m for m in db.query(KnecMark).filter_by(
                subject_id=subject_id, semester_id=semester_id).all()}

        labels_for_template = list(KASNEB_LABELS) if atype == "kasneb" else get_periodic_labels(course)

        return render_template(
            "admin/marks_entry.html",
            subject=subject, semester=semester, course=course,
            department=department, students=students, existing=existing,
            assessment_type=atype,
            periodic_labels=labels_for_template,
        )
    finally:
        db.close()


@app.route("/admin/modules", methods=["GET"])
def admin_modules_picker():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        courses = db.query(Course).order_by(Course.name.asc()).all()
        course_rows = []
        for c in courses:
            course_rows.append({
                "course": c,
                "modules_count": db.query(Module).filter_by(course_id=c.id).count(),
                "subjects_count": db.query(Subject).filter_by(course_id=c.id).count(),
                "students_count": db.query(StudentProfile).filter_by(course_id=c.id).count(),
            })
        return render_template(
            "admin/modules_picker.html",
            course_rows=course_rows, username=session.get("username"),
        )
    finally:
        db.close()


@app.route("/admin/marks/class", methods=["GET"])
def admin_class_marks():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    semester_id = request.args.get("semester_id", type=int)
    if not semester_id:
        flash("Pick a semester.", "warning")
        return redirect(url_for("admin_marks_picker"))

    db = SessionLocal()
    try:
        semester = semester_query(db).filter_by(id=semester_id).first()
        if not semester:
            return "Semester not found", 404

        course = semester.course
        atype = get_course_assessment_type(course)

        students = (
            student_query(db).filter_by(course_id=semester.course_id)
            .order_by(StudentProfile.full_name.asc()).all()
        )
        subjects = (
            subject_query(db).filter_by(course_id=semester.course_id, semester_id=semester_id)
            .order_by(Subject.name.asc()).all()
        )

        if atype == "short":
            marks = {}
            for m in db.query(ShortCourseMark).filter_by(semester_id=semester_id).all():
                marks.setdefault(m.student_id, {})[m.subject_id] = m
        elif atype in ("kasneb", "periodic"):
            marks = {}
            for m in db.query(PeriodicMark).filter_by(semester_id=semester_id).all():
                marks.setdefault(m.student_id, {}).setdefault(m.subject_id, {})[m.label] = m
        else:
            marks = {}
            for m in db.query(KnecMark).filter_by(semester_id=semester_id).all():
                marks.setdefault(m.student_id, {})[m.subject_id] = m

        labels = list(KASNEB_LABELS) if atype == "kasneb" else get_periodic_labels(course)

        return render_template(
            "admin/class_marks.html",
            semester=semester, students=students, subjects=subjects, marks=marks,
            assessment_type=atype,
            periodic_labels=labels,
        )
    finally:
        db.close()


@app.route("/admin/marks/class/export")
def admin_class_marks_export():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    semester_id = request.args.get("semester_id", type=int)
    if not semester_id:
        flash("Pick a semester first.", "warning")
        return redirect(url_for("admin_marks_picker"))

    db = SessionLocal()
    try:
        semester = semester_query(db).filter_by(id=semester_id).first()
        if not semester:
            return "Semester not found", 404

        course = semester.course
        atype = get_course_assessment_type(course)

        students = (
            student_query(db).filter_by(course_id=semester.course_id)
            .order_by(StudentProfile.full_name.asc()).all()
        )
        subjects = (
            subject_query(db).filter_by(course_id=semester.course_id, semester_id=semester_id)
            .order_by(Subject.name.asc()).all()
        )

        data = []
        if atype == "knec":
            marks = {}
            for m in db.query(KnecMark).filter_by(semester_id=semester_id).all():
                marks.setdefault(m.student_id, {})[m.subject_id] = m
            for s in students:
                row = {"Admission No": s.admission_number, "Student Name": s.full_name,
                       "Exam Type": (s.exam_type or "").upper()}
                grand_total = counted = 0
                for subj in subjects:
                    m = marks.get(s.id, {}).get(subj.id)
                    c1 = m.cat1 if m and m.cat1 is not None else ""
                    c2 = m.cat2 if m and m.cat2 is not None else ""
                    fn = m.final if m and m.final is not None else ""
                    total = (m.cat1 or 0) + (m.cat2 or 0) + (m.final or 0) if m else ""
                    prefix = subj.name.split(" ", 1)[0]
                    row[f"{prefix} CAT1"] = c1
                    row[f"{prefix} CAT2"] = c2
                    row[f"{prefix} Final"] = fn
                    row[f"{prefix} Total"] = total
                    if isinstance(total, int):
                        grand_total += total
                        counted += 1
                row["Grand Total"] = grand_total
                row["Average"] = round(grand_total / counted, 2) if counted else ""
                data.append(row)

        elif atype == "short":
            marks = {}
            for m in db.query(ShortCourseMark).filter_by(semester_id=semester_id).all():
                marks.setdefault(m.student_id, {})[m.subject_id] = m
            for s in students:
                row = {"Admission No": s.admission_number, "Student Name": s.full_name}
                grand_total = counted = 0
                for subj in subjects:
                    m = marks.get(s.id, {}).get(subj.id)
                    score = m.score if m and m.score is not None else ""
                    row[f"{subj.name} Score"] = score
                    row[f"{subj.name} Remarks"] = m.remarks if m and m.remarks else ""
                    if isinstance(score, (int, float)):
                        grand_total += score
                        counted += 1
                row["Grand Total"] = grand_total
                row["Average"] = round(grand_total / counted, 2) if counted else ""
                data.append(row)

        elif atype == "kasneb":
            marks = {}
            for m in db.query(PeriodicMark).filter_by(semester_id=semester_id).all():
                marks.setdefault(m.student_id, {}).setdefault(m.subject_id, {})[m.label] = m
            for s in students:
                row = {"Admission No": s.admission_number, "Student Name": s.full_name}
                for subj in subjects:
                    cells = marks.get(s.id, {}).get(subj.id, {})
                    row[f"{subj.name} CAT1"] = cells["CAT 1"].score if "CAT 1" in cells else ""
                    row[f"{subj.name} CAT2"] = cells["CAT 2"].score if "CAT 2" in cells else ""
                data.append(row)

        else:  # periodic
            labels = get_periodic_labels(course)
            marks = {}
            for m in db.query(PeriodicMark).filter_by(semester_id=semester_id).all():
                marks.setdefault(m.student_id, {}).setdefault(m.subject_id, {})[m.label] = m
            for s in students:
                row = {"Admission No": s.admission_number, "Student Name": s.full_name}
                for subj in subjects:
                    cells = marks.get(s.id, {}).get(subj.id, {})
                    for label in labels:
                        row[f"{subj.name} {label}"] = cells[label].score if label in cells else ""
                data.append(row)

        df = pd.DataFrame(data)
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Marks")
        buf.seek(0)

        safe_course = (course.name if course else "Class").replace(" ", "_")[:30]
        safe_sem = semester.name.replace(" ", "_")

        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"Marks_{safe_course}_{safe_sem}.xlsx",
        )
    finally:
        db.close()


@app.route("/admin/reports", methods=["GET"])
def admin_reports_picker():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        courses = db.query(Course).order_by(Course.name.asc()).all()
        semesters = (
            semester_query(db)
            .order_by(Semester.course_id.asc(), Semester.created_at.desc())
            .all()
        )
        return render_template("admin/reports_picker.html", courses=courses, semesters=semesters)
    finally:
        db.close()


@app.route("/admin/reports/class", methods=["GET"])
def admin_reports_class():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    semester_id = request.args.get("semester_id", type=int)
    if not semester_id:
        flash("Pick a semester first.", "warning")
        return redirect(url_for("admin_reports_picker"))

    db = SessionLocal()
    try:
        semester = semester_query(db).filter_by(id=semester_id).first()
        if not semester:
            return "Semester not found", 404

        course = semester.course
        atype = get_course_assessment_type(course)

        students = (
            student_query(db).filter_by(course_id=semester.course_id)
            .order_by(StudentProfile.full_name.asc()).all()
        )

        student_summary = []
        for s in students:
            total_sum = 0
            subjects_count = 0

            if atype == "knec":
                for m in db.query(KnecMark).filter_by(
                    student_id=s.id, semester_id=semester_id
                ).all():
                    total_sum += (m.cat1 or 0) + (m.cat2 or 0) + (m.final or 0)
                    subjects_count += 1
            elif atype == "short":
                for m in db.query(ShortCourseMark).filter_by(
                    student_id=s.id, semester_id=semester_id
                ).all():
                    if m.score is not None:
                        total_sum += m.score
                        subjects_count += 1
            else:
                rows = db.query(PeriodicMark).filter_by(
                    student_id=s.id, semester_id=semester_id
                ).all()
                by_subject = {}
                for r in rows:
                    by_subject.setdefault(r.subject_id, []).append(r.score)
                for scores in by_subject.values():
                    valid = [sc for sc in scores if sc is not None]
                    if valid:
                        total_sum += sum(valid) / len(valid)
                        subjects_count += 1

            avg = round(total_sum / subjects_count, 2) if subjects_count else None

            if avg is None:
                grade = result = None
            elif atype == "knec":
                grade = (1 if avg >= 84 else 2 if avg >= 75 else 3 if avg >= 66
                         else 4 if avg >= 60 else 5 if avg >= 50
                         else 6 if avg >= 40 else 7)
                result = ("DISTINCTION" if avg >= 75 else "CREDIT" if avg >= 60
                          else "PASS" if avg >= 40 else "REFER")
            else:
                grade = (1 if avg >= 80 else 2 if avg >= 65 else 3 if avg >= 50 else 4)
                result = ("DISTINCTION" if avg >= 80 else "CREDIT" if avg >= 65
                          else "PASS" if avg >= 50 else "FAIL")

            student_summary.append({
                "student": s, "avg": avg,
                "overall_grade": grade, "overall_result": result,
                "subjects_count": subjects_count,
            })

        return render_template(
            "admin/reports_class.html",
            semester=semester, student_summary=student_summary,
            total_students=len(students),
            assessment_type=atype,
        )
    finally:
        db.close()


@app.route("/admin/reports/class/export")
def admin_reports_class_export():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    semester_id = request.args.get("semester_id", type=int)
    if not semester_id:
        flash("Pick a semester first.", "warning")
        return redirect(url_for("admin_reports_picker"))

    db = SessionLocal()
    try:
        semester = semester_query(db).filter_by(id=semester_id).first()
        if not semester:
            return "Semester not found", 404

        course = semester.course
        atype = get_course_assessment_type(course)

        students = (
            student_query(db).filter_by(course_id=semester.course_id)
            .order_by(StudentProfile.full_name.asc()).all()
        )

        data = []
        for s in students:
            total_sum = 0
            subjects_count = 0

            if atype == "knec":
                for m in db.query(KnecMark).filter_by(student_id=s.id, semester_id=semester_id).all():
                    total_sum += (m.cat1 or 0) + (m.cat2 or 0) + (m.final or 0)
                    subjects_count += 1
            elif atype == "short":
                for m in db.query(ShortCourseMark).filter_by(student_id=s.id, semester_id=semester_id).all():
                    if m.score is not None:
                        total_sum += m.score
                        subjects_count += 1
            else:
                rows = db.query(PeriodicMark).filter_by(student_id=s.id, semester_id=semester_id).all()
                by_subject = {}
                for r in rows:
                    by_subject.setdefault(r.subject_id, []).append(r.score)
                for scores in by_subject.values():
                    valid = [sc for sc in scores if sc is not None]
                    if valid:
                        total_sum += sum(valid) / len(valid)
                        subjects_count += 1

            avg = round(total_sum / subjects_count, 2) if subjects_count else ""

            if not isinstance(avg, float):
                grade = result = ""
            elif atype == "knec":
                grade = (1 if avg >= 84 else 2 if avg >= 75 else 3 if avg >= 66
                         else 4 if avg >= 60 else 5 if avg >= 50
                         else 6 if avg >= 40 else 7)
                result = ("DISTINCTION" if avg >= 75 else "CREDIT" if avg >= 60
                          else "PASS" if avg >= 40 else "REFER")
            else:
                grade = (1 if avg >= 80 else 2 if avg >= 65 else 3 if avg >= 50 else 4)
                result = ("DISTINCTION" if avg >= 80 else "CREDIT" if avg >= 65
                          else "PASS" if avg >= 50 else "FAIL")

            data.append({
                "Admission No": s.admission_number,
                "Student Name": s.full_name,
                "Course": course.name if course else "",
                "Semester": semester.name,
                "Exam Type": (s.exam_type or "").upper(),
                "Assessment Type": atype.upper(),
                "Subjects Graded": subjects_count,
                "Total Marks": total_sum,
                "Average": avg, "Grade": grade, "Overall Result": result,
            })

        df = pd.DataFrame(data)
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Class Report")
        buf.seek(0)

        safe_course = (course.name if course else "Class").replace(" ", "_")[:30]
        safe_sem = semester.name.replace(" ", "_")

        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"ClassReport_{safe_course}_{safe_sem}.xlsx",
        )
    finally:
        db.close()


@app.route("/admin/analytics/knec/<int:course_id>/export")
def admin_knec_analytics_export(course_id):
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        course = db.query(Course).filter_by(id=course_id).first()
        if not course or get_course_assessment_type(course) != "knec":
            flash("KNEC course not found.", "danger")
            return redirect(url_for("admin_reports_picker"))

        buf = _build_knec_analytics_workbook(db, course)

        safe_course = course.name.replace(" ", "_")[:30]
        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"KNEC_Analytics_{safe_course}.xlsx",
        )
    finally:
        db.close()


@app.route("/admin/analytics/knec", methods=["GET"])
def admin_knec_analytics_picker():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        courses = (
            db.query(Course)
            .order_by(Course.name.asc())
            .all()
        )
        knec_courses = [
            c for c in courses
            if (c.assessment_type or "knec").strip().lower() == "knec"
        ]

        course_rows = []
        for c in knec_courses:
            students_count = db.query(StudentProfile).filter_by(course_id=c.id).count()
            semesters_count = db.query(Semester).filter_by(course_id=c.id).count()
            subjects_count = db.query(Subject).filter_by(course_id=c.id).count()
            marks_count = (
                db.query(KnecMark)
                .join(StudentProfile, KnecMark.student_id == StudentProfile.id)
                .filter(StudentProfile.course_id == c.id)
                .count()
            )
            course_rows.append({
                "course": c,
                "students_count": students_count,
                "semesters_count": semesters_count,
                "subjects_count": subjects_count,
                "marks_count": marks_count,
            })

        return render_template(
            "admin/knec_analytics_picker.html",
            course_rows=course_rows,
            username=session.get("username"),
        )
    finally:
        db.close()


# ═════════════════════════════════════════════════════════════
# ADMIN FILTERED REPORT
# ═════════════════════════════════════════════════════════════
@app.route("/admin/reports/filtered_export", methods=["GET", "POST"])
def admin_filtered_report_export():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        course_id   = request.args.get("course_id", type=int)
        semester_id = request.args.get("semester_id", type=int)
        module_id   = request.args.get("module_id", type=int)
        export      = request.args.get("export")

        courses   = db.query(Course).order_by(Course.name.asc()).all()
        semesters = (
            semester_query(db)
            .order_by(Semester.course_id.asc(), Semester.created_at.asc())
            .all()
        )
        modules = (
            db.query(Module)
            .order_by(Module.course_id.asc(), Module.order.asc(), Module.name.asc())
            .all()
        )

        selected_course   = None
        selected_semester = None
        selected_module   = None
        report_rows       = []

        if course_id:
            selected_course = db.query(Course).filter_by(id=course_id).first()
        if semester_id:
            selected_semester = semester_query(db).filter_by(id=semester_id).first()
        if module_id:
            selected_module = db.query(Module).filter_by(id=module_id).first()

        if selected_course and selected_semester:
            report_rows = _build_filtered_report_rows(
                db, selected_course, selected_semester, selected_module
            )

        if export == "true":
            if not report_rows:
                flash("No data to export for that selection.", "warning")
                return redirect(url_for("admin_filtered_report_export",
                                        course_id=course_id,
                                        semester_id=semester_id,
                                        module_id=module_id))
            return _send_filtered_report_excel(
                report_rows, selected_course, selected_semester, selected_module,
                header_color="1ABC9C", prefix="Report"
            )

        return render_template(
            "admin/filtered_report.html",
            courses=courses,
            semesters=semesters,
            modules=modules,
            selected_course=selected_course,
            selected_semester=selected_semester,
            selected_module=selected_module,
            report_rows=report_rows,
            username=session.get("username"),
        )
    finally:
        db.close()


@app.route("/admin/reports/student/<int:student_id>/<int:semester_id>")
def admin_report_student(student_id, semester_id):
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        student = student_query(db).filter_by(id=student_id).first()
        if not student:
            return "Student not found", 404

        sem = db.query(Semester).filter_by(id=semester_id, course_id=student.course_id).first()
        if not sem:
            return "Semester not found for this student", 404

        course = student.course
        atype = get_course_assessment_type(course)

        subjects = (
            subject_query(db).filter_by(course_id=student.course_id, semester_id=sem.id)
            .order_by(Subject.name.asc()).all()
        )

        # ── Only include subjects the student actually has marks for ──
        rows = []
        grand_total = 0
        counted = 0

        for subj in subjects:
            info = get_student_subject_marks(db, student.id, subj.id, sem.id, course)
            if not info:
                continue
            rows.append({"subject": subj, "info": info})
            grand_total += info["total"]
            counted += 1

        avg = round(grand_total / counted, 2) if counted else None

        if not rows:
            flash(
                "No marks have been entered for this student in the selected semester yet.",
                "warning",
            )
            return redirect(url_for("admin_reports_picker"))

        labels = list(KASNEB_LABELS) if atype == "kasneb" else get_periodic_labels(course)

        return render_template(
            "students/print_report.html",
            student=student, semester=sem, rows=rows,
            assessment_type=atype,
            periodic_labels=labels,
            grand_total=grand_total, counted=counted, avg=avg,
            current_year=datetime.now().year, admin_mode=True,
        )
    finally:
        db.close()

"""
============================================================
SHARED FILTERED REPORT HELPERS
============================================================
"""
def _build_filtered_report_rows(db, course, semester, module_obj=None):
    """Build the list of dict rows for the filtered report (any assessment type)."""
    atype = get_course_assessment_type(course)

    students_q = student_query(db).filter_by(course_id=course.id)
    if module_obj:
        students_q = students_q.filter(
            or_(
                StudentProfile.module_id == module_obj.id,
                StudentProfile.module_id.is_(None),
            )
        )
    students = students_q.order_by(StudentProfile.full_name.asc()).all()

    subjects = (
        subject_query(db)
        .filter_by(course_id=course.id, semester_id=semester.id)
        .order_by(Subject.name.asc())
        .all()
    )

    student_ids = [s.id for s in students]

    if atype == "knec":
        raw = (
            db.query(KnecMark)
            .filter(KnecMark.student_id.in_(student_ids),
                    KnecMark.semester_id == semester.id)
            .all()
        ) if student_ids else []
        marks_by_key = {(m.student_id, m.subject_id): m for m in raw}
    elif atype == "short":
        raw = (
            db.query(ShortCourseMark)
            .filter(ShortCourseMark.student_id.in_(student_ids),
                    ShortCourseMark.semester_id == semester.id)
            .all()
        ) if student_ids else []
        marks_by_key = {(m.student_id, m.subject_id): m for m in raw}
    else:
        raw = (
            db.query(PeriodicMark)
            .filter(PeriodicMark.student_id.in_(student_ids),
                    PeriodicMark.semester_id == semester.id)
            .all()
        ) if student_ids else []
        marks_by_key = {}
        for m in raw:
            marks_by_key.setdefault((m.student_id, m.subject_id), {})[m.label] = m

    labels = (
        list(KASNEB_LABELS) if atype == "kasneb"
        else get_periodic_labels(course)
    )

    report_rows = []

    for s in students:
        student_total = 0.0
        student_counted = 0

        for subj in subjects:
            if atype == "knec":
                m = marks_by_key.get((s.id, subj.id))
                cat1 = m.cat1 if m else None
                cat2 = m.cat2 if m else None
                final = m.final if m else None
                total = (cat1 or 0) + (cat2 or 0) + (final or 0) if m else ""
                result = ""
                if isinstance(total, (int, float)) and total > 0:
                    result = knec_result(total)
                    student_total += total
                    student_counted += 1
                report_rows.append({
                    "Admission No": s.admission_number,
                    "Student Name": s.full_name,
                    "Course": course.name,
                    "Module": s.module or "—",
                    "Semester": semester.name,
                    "Subject": subj.name,
                    "CAT 1": cat1 if cat1 is not None else "",
                    "CAT 2": cat2 if cat2 is not None else "",
                    "Final": final if final is not None else "",
                    "Total": total,
                    "Result": result,
                })

            elif atype == "short":
                m = marks_by_key.get((s.id, subj.id))
                score = m.score if m else None
                remarks = m.remarks if m and m.remarks else ""
                result = ""
                if isinstance(score, (int, float)):
                    result = (
                        "DISTINCTION" if score >= 80 else
                        "CREDIT"      if score >= 65 else
                        "PASS"        if score >= 50 else "FAIL"
                    )
                    student_total += score
                    student_counted += 1
                report_rows.append({
                    "Admission No": s.admission_number,
                    "Student Name": s.full_name,
                    "Course": course.name,
                    "Module": s.module or "—",
                    "Semester": semester.name,
                    "Subject": subj.name,
                    "Score": score if score is not None else "",
                    "Remarks": remarks,
                    "Result": result,
                })

            else:  # kasneb / periodic
                cells = marks_by_key.get((s.id, subj.id), {})
                row = {
                    "Admission No": s.admission_number,
                    "Student Name": s.full_name,
                    "Course": course.name,
                    "Module": s.module or "—",
                    "Semester": semester.name,
                    "Subject": subj.name,
                }
                scores = []
                for label in labels:
                    m = cells.get(label)
                    val = m.score if m else ""
                    row[label] = val
                    if isinstance(val, (int, float)):
                        scores.append(val)
                avg = round(sum(scores) / len(scores), 2) if scores else ""
                row["Average"] = avg
                row["Result"] = (
                    ("DISTINCTION" if avg >= 80 else
                     "CREDIT"      if avg >= 65 else
                     "PASS"        if avg >= 50 else
                     "FAIL") if isinstance(avg, float) else ""
                )
                if isinstance(avg, float):
                    student_total += avg
                    student_counted += 1
                report_rows.append(row)

        if student_counted:
            summary = {
                "Admission No": s.admission_number,
                "Student Name": s.full_name,
                "Course": course.name,
                "Module": s.module or "—",
                "Semester": semester.name,
                "Subject": "— OVERALL AVERAGE —",
            }
            if atype == "knec":
                summary["CAT 1"] = ""
                summary["CAT 2"] = ""
                summary["Final"] = ""
                summary["Total"] = round(student_total, 2)
                summary["Result"] = knec_result(
                    round(student_total / student_counted, 2)
                )
            elif atype == "short":
                summary["Score"] = round(student_total, 2)
                summary["Remarks"] = ""
                summary["Result"] = (
                    "DISTINCTION" if (student_total / student_counted) >= 80 else
                    "CREDIT"      if (student_total / student_counted) >= 65 else
                    "PASS"        if (student_total / student_counted) >= 50 else "FAIL"
                )
            else:
                avg = round(student_total / student_counted, 2)
                for label in labels:
                    summary[label] = ""
                summary["Average"] = avg
                summary["Result"] = (
                    "DISTINCTION" if avg >= 80 else
                    "CREDIT"      if avg >= 65 else
                    "PASS"        if avg >= 50 else "FAIL"
                )
            report_rows.append(summary)

    return report_rows


def _send_filtered_report_excel(report_rows, course, semester, module_obj,
                                header_color="1ABC9C", prefix="Report"):
    """Convert rows to Excel and return the Flask send_file response."""
    df = pd.DataFrame(report_rows)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Report")
        ws = writer.sheets["Report"]

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor=header_color)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col) + 3
            ws.column_dimensions[col[0].column_letter].width = min(width, 40)

        ws.freeze_panes = "A2"

    buf.seek(0)

    parts = [course.name.replace(" ", "_")[:30], semester.name.replace(" ", "_")]
    if module_obj:
        parts.append(module_obj.name.replace(" ", "_"))
    fname = prefix + "_" + "_".join(parts) + ".xlsx"

    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=fname,
    )


"""
============================================================HOD ROUTES
============================================================
"""
@app.route("/hod/dashboard")
def hod_dashboard():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        department = db.query(Department).filter_by(id=session["department_id"]).first()
        if not department:
            flash("Department not found.", "danger")
            return redirect(url_for("logout"))

        courses = (
            db.query(Course)
            .filter_by(department_id=department.id)
            .order_by(Course.name.asc())
            .all()
        )
        course_ids = [c.id for c in courses]

        students_count = subjects_count = semesters_count = 0
        if course_ids:
            students_count = (
                db.query(func.count(StudentProfile.id))
                .filter(StudentProfile.course_id.in_(course_ids))
                .scalar()
            ) or 0
            subjects_count = (
                db.query(func.count(Subject.id))
                .filter(Subject.course_id.in_(course_ids))
                .scalar()
            ) or 0
            semesters_count = (
                db.query(func.count(Semester.id))
                .filter(Semester.course_id.in_(course_ids))
                .scalar()
            ) or 0

        STUDENTS_PER_PAGE = 8
        student_page = request.args.get("student_page", 1, type=int) or 1
        if student_page < 1:
            student_page = 1

        student_search = request.args.get("student_search", "").strip()

        students_pagination = None
        students = []
        no_sem_count = 0

        if course_ids:
            base_q = student_query(db).filter(StudentProfile.course_id.in_(course_ids))
            if student_search:
                like = f"%{student_search}%"
                base_q = base_q.filter(or_(
                    StudentProfile.full_name.ilike(like),
                    StudentProfile.admission_number.ilike(like),
                    StudentProfile.phone_number.ilike(like),
                ))
            base_q = base_q.order_by(StudentProfile.full_name.asc())
            students_pagination = paginate_query(base_q, student_page, STUDENTS_PER_PAGE)
            students = students_pagination.items

            no_sem_count = (
                db.query(func.count(StudentProfile.id))
                .filter(
                    StudentProfile.course_id.in_(course_ids),
                    StudentProfile.current_semester_id.is_(None),
                )
                .scalar()
            ) or 0

        COURSES_PER_PAGE = 9
        course_page = request.args.get("course_page", 1, type=int) or 1
        if course_page < 1:
            course_page = 1

        course_search = request.args.get("course_search", "").strip()
        course_module_filter = request.args.get("course_module", "").strip()

        card_subjects = {}
        card_current_sem = {}
        card_students_by_module = defaultdict(lambda: defaultdict(int))
        students_by_course = defaultdict(list)
        all_module_labels = set()

        if course_ids:
            all_dept_students = (
                student_query(db)
                .filter(StudentProfile.course_id.in_(course_ids))
                .order_by(StudentProfile.full_name.asc())
                .all()
            )
            for s in all_dept_students:
                label = (s.module or "Unassigned").strip().upper() or "Unassigned"
                students_by_course[s.course_id].append(s)
                card_students_by_module[s.course_id][label] += 1
                all_module_labels.add(label)

            for cid, cnt in (
                db.query(Subject.course_id, func.count(Subject.id))
                .filter(Subject.course_id.in_(course_ids))
                .group_by(Subject.course_id)
                .all()
            ):
                card_subjects[cid] = cnt

            for sem in (
                db.query(Semester)
                .filter(Semester.course_id.in_(course_ids), Semester.is_current == True)
                .all()
            ):
                card_current_sem[sem.course_id] = sem

        def _module_sort_key(label):
            import re as _re
            m = _re.match(r"\s*MODULE\s+(\d+)", str(label), _re.IGNORECASE)
            if m:
                return (0, int(m.group(1)), "")
            return (1, 0, str(label).lower())

        module_options = sorted(all_module_labels, key=_module_sort_key)

        filtered_courses = []
        for c in courses:
            if course_module_filter:
                labels_for_this_course = set(card_students_by_module.get(c.id, {}).keys())
                if course_module_filter not in labels_for_this_course:
                    continue
            if course_search:
                needle = course_search.lower()
                haystack = f"{c.name} {c.level or ''}".lower()
                if needle not in haystack:
                    continue
            filtered_courses.append(c)

        total_filtered = len(filtered_courses)
        total_course_pages = (total_filtered + COURSES_PER_PAGE - 1) // COURSES_PER_PAGE
        if total_course_pages > 0 and course_page > total_course_pages:
            course_page = total_course_pages

        start = (course_page - 1) * COURSES_PER_PAGE
        page_courses = filtered_courses[start:start + COURSES_PER_PAGE]

        course_cards = []
        for c in page_courses:
            modules_breakdown = dict(card_students_by_module.get(c.id, {}))
            total_students = sum(modules_breakdown.values())
            course_cards.append({
                "course": c,
                "current_semester": card_current_sem.get(c.id),
                "students_count": total_students,
                "students_by_module": modules_breakdown,
                "students_list": students_by_course.get(c.id, []),
                "subjects_count": card_subjects.get(c.id, 0),
            })

        course_pagination = {
            "page": course_page,
            "per_page": COURSES_PER_PAGE,
            "total": total_filtered,
            "pages": total_course_pages,
            "has_prev": course_page > 1,
            "has_next": course_page < total_course_pages,
            "prev_num": course_page - 1,
            "next_num": course_page + 1,
            "first": start + 1 if total_filtered else 0,
            "last": min(start + COURSES_PER_PAGE, total_filtered),
        }

        semester_map = {cid: [] for cid in course_ids}
        if course_ids:
            sem_rows = (
                db.query(Semester)
                .filter(Semester.course_id.in_(course_ids))
                .order_by(Semester.created_at.asc())
                .all()
            )
            for sem in sem_rows:
                semester_map[sem.course_id].append({"id": sem.id, "name": sem.name})

        return render_template(
            "hod/dashboard.html",
            hod_name=session.get("username"), department=department,
            course_cards=course_cards,
            course_pagination=course_pagination,
            course_search=course_search,
            course_module_filter=course_module_filter,
            module_options=module_options,
            students=students,
            students_pagination=students_pagination,
            student_search=student_search,
            students_count=students_count, subjects_count=subjects_count,
            semesters_count=semesters_count, no_sem_count=no_sem_count,
            semester_map=semester_map, current_year=datetime.now().year,
        )
    finally:
        db.close()


@app.route("/hod/subjects", methods=["GET"])
def hod_manage_subjects():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    db = SessionLocal()
    try:
        department = db.query(Department).filter_by(id=department_id).first()
        if not department:
            flash("Department not found.", "danger")
            return redirect(url_for("logout"))

        courses = db.query(Course).filter_by(department_id=department_id).order_by(Course.name.asc()).all()
        course_ids = [c.id for c in courses]

        selected_course_id = request.args.get("course_id", type=int)

        subjects_query = subject_query(db).filter(Subject.course_id.in_(course_ids))
        if selected_course_id and selected_course_id in course_ids:
            subjects_query = subjects_query.filter(Subject.course_id == selected_course_id)

        subjects = subjects_query.order_by(Subject.name.asc()).all()

        semesters = []
        if course_ids:
            semesters = (
                semester_query(db).filter(Semester.course_id.in_(course_ids))
                .order_by(Semester.name.asc()).all()
            )
        semesters_json = [
            {"id": s.id, "name": s.name, "course_id": s.course_id}
            for s in semesters
        ]

        return render_template(
            "hod/manage_subjects.html",
            department=department, courses=courses, subjects=subjects,
            semesters=semesters_json, selected_course_id=selected_course_id,
            hod_name=session.get("username"),
        )
    finally:
        db.close()


@app.route("/hod/subjects/add", methods=["POST"])
def hod_add_subject():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    subject_name = request.form.get("subject_name", "").strip()
    course_id = request.form.get("course_id", type=int)
    semester_id = request.form.get("semester_id", type=int)
    module_id = request.form.get("module_id", type=int)

    if not subject_name or not course_id:
        flash("Subject name and course are required.", "danger")
        return redirect(url_for("hod_manage_subjects"))

    db = SessionLocal()
    try:
        course = db.query(Course).filter_by(id=course_id, department_id=department_id).first()
        if not course:
            flash("Course not found in your department.", "danger")
            return redirect(url_for("hod_manage_subjects"))

        if semester_id:
            sem = db.query(Semester).filter_by(id=semester_id).first()
            if not sem or sem.course_id != course_id:
                flash("Selected semester doesn't belong to that course.", "danger")
                return redirect(url_for("hod_manage_subjects"))

        q = db.query(Subject).filter_by(name=subject_name, course_id=course_id)
        if semester_id:
            q = q.filter_by(semester_id=semester_id)
        else:
            q = q.filter(Subject.semester_id.is_(None))

        if q.first():
            flash("Subject already exists for this course/semester.", "warning")
            return redirect(url_for("hod_manage_subjects"))

        module_obj = None
        if module_id:
            module_obj = db.query(Module).filter_by(id=module_id, course_id=course_id).first()

        db.add(Subject(
            name=subject_name, course_id=course_id,
            semester_id=semester_id,
            module=module_obj.name if module_obj else None,
            module_id=module_obj.id if module_obj else None,
        ))
        db.commit()
        flash(f"✅ Subject '{subject_name}' added.", "success")
    except Exception as e:
        db.rollback()
        flash(f"Error: {str(e)}", "danger")
    finally:
        db.close()

    return redirect(url_for("hod_manage_subjects"))


@app.route("/hod/subjects/edit/<int:subject_id>", methods=["GET", "POST"])
def hod_edit_subject(subject_id):
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    db = SessionLocal()
    try:
        subject = subject_query(db).filter_by(id=subject_id).first()
        if not subject:
            flash("Subject not found.", "danger")
            return redirect(url_for("hod_manage_subjects"))

        course = db.query(Course).filter_by(id=subject.course_id).first()
        if not course or course.department_id != department_id:
            flash("Subject not in your department.", "danger")
            return redirect(url_for("hod_manage_subjects"))

        if request.method == "POST":
            name = request.form.get("subject_name", "").strip()
            semester_id = request.form.get("semester_id", type=int)
            module_id = request.form.get("module_id", type=int)

            if not name:
                flash("Subject name is required.", "danger")
                return redirect(url_for("hod_edit_subject", subject_id=subject_id))

            if semester_id:
                sem = db.query(Semester).filter_by(id=semester_id).first()
                if not sem or sem.course_id != subject.course_id:
                    flash("Selected semester doesn't belong to that course.", "danger")
                    return redirect(url_for("hod_edit_subject", subject_id=subject_id))

            subject.name = name
            subject.semester_id = semester_id

            if module_id:
                module_obj = db.query(Module).filter_by(id=module_id, course_id=subject.course_id).first()
                if module_obj:
                    subject.module_id = module_obj.id
                    subject.module = module_obj.name
            else:
                subject.module_id = None
                subject.module = None

            db.commit()
            flash("✅ Subject updated.", "success")
            return redirect(url_for("hod_manage_subjects"))

        semesters = (
            semester_query(db).filter_by(course_id=subject.course_id)
            .order_by(Semester.name.asc()).all()
        )
        modules = (
            db.query(Module).filter_by(course_id=subject.course_id)
            .order_by(Module.order.asc(), Module.name.asc()).all()
        )
        return render_template(
            "hod/edit_subject.html",
            subject=subject, semesters=semesters, modules=modules,
            department=db.query(Department).filter_by(id=department_id).first(),
            hod_name=session.get("username"),
        )
    finally:
        db.close()


@app.route("/hod/subjects/delete/<int:subject_id>", methods=["POST"])
def hod_delete_subject(subject_id):
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    db = SessionLocal()
    try:
        subject = db.query(Subject).filter_by(id=subject_id).first()
        if not subject:
            flash("Subject not found.", "danger")
            return redirect(url_for("hod_manage_subjects"))

        course = db.query(Course).filter_by(id=subject.course_id).first()
        if not course or course.department_id != department_id:
            flash("Subject not in your department.", "danger")
            return redirect(url_for("hod_manage_subjects"))

        _delete_quizzes_for(db, subject_id=subject_id)
        db.query(Message).filter(Message.subject_id == subject_id).delete(synchronize_session=False)
        db.query(KnecMark).filter(KnecMark.subject_id == subject_id).delete(synchronize_session=False)
        db.query(ShortCourseMark).filter(ShortCourseMark.subject_id == subject_id).delete(synchronize_session=False)
        db.query(PeriodicMark).filter(PeriodicMark.subject_id == subject_id).delete(synchronize_session=False)
        db.delete(subject)
        db.commit()
        flash("✅ Subject deleted.", "success")
    except Exception as e:
        db.rollback()
        flash(f"Error: {str(e)}", "danger")
    finally:
        db.close()

    return redirect(url_for("hod_manage_subjects"))


@app.route("/hod/students", methods=["GET", "POST"])
def hod_students_list():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    db = SessionLocal()
    try:
        department = db.query(Department).filter_by(id=department_id).first()
        if not department:
            flash("Department not found.", "danger")
            return redirect(url_for("logout"))

        courses = db.query(Course).filter_by(department_id=department_id).order_by(Course.name.asc()).all()
        course_ids = [c.id for c in courses]

        if request.method == "POST":
            student_id = request.form.get("student_id", type=int)
            action = request.form.get("action")

            student = db.query(StudentProfile).filter_by(id=student_id).first()
            if not student or student.course_id not in course_ids:
                flash("Student not found in your department.", "danger")
                return redirect(url_for("hod_students_list"))

            if action == "toggle_block":
                student.blocked = not student.blocked
                db.commit()
                state = "blocked" if student.blocked else "activated"
                flash(f"✅ Student {state}.", "success")
            elif action == "delete":
                try:
                    _delete_student_tree(db, student)
                    db.commit()
                    flash("✅ Student deleted.", "success")
                except Exception as e:
                    db.rollback()
                    flash(f"Error deleting student: {e}", "danger")

            return redirect(url_for("hod_students_list"))

        search_query_str = request.args.get("search", "").strip()

        if course_ids:
            q = student_query(db).filter(StudentProfile.course_id.in_(course_ids))
            if search_query_str:
                q = q.filter(or_(
                    StudentProfile.full_name.ilike(f"%{search_query_str}%"),
                    StudentProfile.admission_number.ilike(f"%{search_query_str}%"),
                    StudentProfile.phone_number.ilike(f"%{search_query_str}%"),
                ))
            students = q.order_by(StudentProfile.full_name.asc()).all()
        else:
            students = []

        semester_map = {}
        for cid in course_ids:
            sems = db.query(Semester).filter_by(course_id=cid).order_by(Semester.created_at.asc()).all()
            semester_map[cid] = [{"id": s.id, "name": s.name} for s in sems]

        return render_template(
            "hod/students_list.html",
            department=department, hod_name=session.get("username"),
            students=students, search_query=search_query_str,
            semester_map=semester_map,
            courses=courses,
        )
    finally:
        db.close()


@app.route("/hod/students/add", methods=["GET", "POST"])
def hod_add_student():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    db = SessionLocal()
    try:
        department = db.query(Department).filter_by(id=department_id).first()
        if not department:
            flash("Department not found.", "danger")
            return redirect(url_for("logout"))

        courses = db.query(Course).filter_by(
            department_id=department_id
        ).order_by(Course.name.asc()).all()

        if request.method == "POST":
            data, errors = _validate_new_student_form(
                db, request.form, department_id=department_id
            )

            if errors:
                for e in errors:
                    flash(f"❌ {e}", "danger")

                # ── C2: repopulate dependent dropdowns for the re-render ──
                error_course_id = request.form.get("course_id", type=int)
                error_semesters = []
                error_modules = []
                if error_course_id:
                    error_semesters = (
                        db.query(Semester)
                        .filter_by(course_id=error_course_id)
                        .order_by(Semester.created_at.asc())
                        .all()
                    )
                    error_modules = (
                        db.query(Module)
                        .filter_by(course_id=error_course_id)
                        .order_by(Module.order.asc(), Module.name.asc())
                        .all()
                    )

                return render_template(
                    "hod/add_student.html",
                    department=department,
                    hod_name=session.get("username"),
                    courses=courses,
                    semesters=error_semesters,
                    modules=error_modules,
                    form=request.form,
                    errors=errors,
                    error_fields=_error_fields_from(errors),
                )

            ok, msg, _ = _create_single_student(
                db,
                full_name=data["full_name"], admission=data["admission"],
                course=data["course"], exam_type=data["exam_type"],
                semester_id=data["semester_id"], module_obj=data["module_obj"],
                phone=data["phone"], username=data["username"],
                password=data["password"],
                strict_department_id=department_id,
            )

            if not ok:
                db.rollback()
                flash(f"❌ {msg}", "danger")
                return render_template(
                    "hod/add_student.html",
                    department=department,
                    hod_name=session.get("username"),
                    courses=courses,
                    semesters=[],
                    modules=[],
                    form=request.form,
                    errors=[msg],
                    error_fields=set(),
                )

            db.commit()
            flash(f"✅ {msg}", "success")
            return redirect(url_for("hod_students_list"))

        # ── GET branch ──
        return render_template(
            "hod/add_student.html",
            department=department,
            hod_name=session.get("username"),
            courses=courses,
            semesters=[],
            modules=[],
            form={},
            errors=[],
            error_fields=set(),
        )

    except IntegrityError as e:
        db.rollback()
        traceback.print_exc()
        flash(f"❌ {_translate_db_error(e)}", "danger")
        return redirect(url_for("hod_add_student"))

    except Exception:
        db.rollback()
        traceback.print_exc()
        flash("❌ An unexpected error occurred. Please try again.", "danger")
        return redirect(url_for("hod_add_student"))

    finally:
        db.close()

@app.route("/hod/students/upload/template")
def hod_students_template():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        department_id = session["department_id"]
        courses = db.query(Course).filter_by(department_id=department_id).order_by(Course.name.asc()).all()

        sample = pd.DataFrame([{
            "full_name": "silas sakwa",
            "admission_number": "STC/2025/008",
            "course_name": courses[0].name if courses else "CERTIFICATE IN ICT",
            "exam_type": "knec",
            "module": "MODULE 1",
            "semester": "OCT-DEC 2025",
            "phone_number": "0745361425",
            "username": "silas",
            "password": "1234",
        }])

        ref_rows = []
        for c in courses:
            sems = db.query(Semester).filter_by(course_id=c.id).order_by(Semester.created_at.asc()).all()
            if sems:
                for s in sems:
                    ref_rows.append({
                        "course_name": c.name, "level": c.level,
                        "semester (term)": s.name,
                    })
            else:
                ref_rows.append({
                    "course_name": c.name, "level": c.level,
                    "semester (term)": "(no semesters yet)",
                })

        reference = pd.DataFrame(
            ref_rows or [{"course_name": "(no courses yet)", "level": "", "semester (term)": ""}]
        )

        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            sample.to_excel(writer, index=False, sheet_name="Students")
            reference.to_excel(writer, index=False, sheet_name="Valid Courses")

            for name, df in (("Students", sample), ("Valid Courses", reference)):
                ws = writer.sheets[name]
                for idx, col in enumerate(df.columns, start=1):
                    max_len = max(
                        [len(str(col))] + [len(str(v)) for v in df[col].tolist()]
                    ) if len(df) else len(str(col))
                    ws.column_dimensions[
                        ws.cell(row=1, column=idx).column_letter
                    ].width = min(max(max_len + 3, 12), 45)

        buf.seek(0)

        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True, download_name="student_upload_template.xlsx",
        )
    finally:
        db.close()


@app.route("/hod/students/upload", methods=["GET", "POST"])
def hod_students_upload():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    db = SessionLocal()
    try:
        department = db.query(Department).filter_by(id=department_id).first()
        if not department:
            flash("Department not found.", "danger")
            return redirect(url_for("logout"))

        allowed_courses = db.query(Course).filter_by(department_id=department_id).all()
        course_by_name = {c.name.strip().lower(): c for c in allowed_courses}

        if request.method == "POST":
            file = request.files.get("student_file")
            if not file or not file.filename.lower().endswith((".xlsx", ".xls")):
                flash("❌ Please upload a valid .xlsx or .xls file.", "danger")
                return redirect(url_for("hod_students_upload"))

            try:
                df = pd.read_excel(file, sheet_name="Students", dtype=str).fillna("")
            except Exception as e:
                flash(f"❌ Could not read the Excel file: {e}", "danger")
                return redirect(url_for("hod_students_upload"))

            df.columns = [str(c).strip().lower().replace(" ", "_") for c in df.columns]

            required_cols = {"full_name", "admission_number", "course_name", "exam_type"}
            missing = required_cols - set(df.columns)
            if missing:
                flash(f"❌ Missing required columns: {', '.join(sorted(missing))}", "danger")
                return redirect(url_for("hod_students_upload"))

            existing_adm = {a[0] for a in db.query(StudentProfile.admission_number).all() if a[0]}
            seen_in_file = set()

            valid_rows, error_rows = [], []

            for idx, row in df.iterrows():
                row_no = idx + 2
                full_name = str(row.get("full_name", "")).strip()
                admission = str(row.get("admission_number", "")).strip()
                course_name = str(row.get("course_name", "")).strip()
                exam_type = str(row.get("exam_type", "")).strip().lower()
                module_name = str(row.get("module", "")).strip()
                semester_name = str(row.get("semester", "")).strip()
                phone = str(row.get("phone_number", "")).strip()
                username = str(row.get("username", "")).strip()
                password = str(row.get("password", "")).strip()

                errors = []
                semester_id = None

                if not full_name:
                    errors.append("full_name is empty")
                if not admission:
                    errors.append("admission_number is empty")
                elif admission in existing_adm:
                    errors.append(f"admission_number '{admission}' already exists")
                elif admission in seen_in_file:
                    errors.append(f"admission_number '{admission}' duplicated inside the file")
                else:
                    seen_in_file.add(admission)

                matched_course = course_by_name.get(course_name.lower())
                if not matched_course:
                    errors.append(f"course '{course_name}' not found in your department")
                else:
                    if get_course_assessment_type(matched_course) in ("short", "periodic"):
                        if exam_type != "internal":
                            exam_type = "internal"

                    if semester_name:
                        sem = db.query(Semester).filter(
                            Semester.course_id == matched_course.id,
                            Semester.name.ilike(semester_name),
                        ).first()
                        if not sem:
                            errors.append(f"semester '{semester_name}' not found for course '{course_name}'")
                        else:
                            semester_id = sem.id

                if exam_type not in ("knec", "internal"):
                    errors.append("exam_type must be 'knec' or 'internal'")

                if errors:
                    error_rows.append({
                        "row": row_no, "full_name": full_name,
                        "admission_number": admission, "errors": errors,
                    })
                else:
                    valid_rows.append({
                        "row": row_no,
                        "full_name": full_name,
                        "admission_number": admission,
                        "course_name": course_name,
                        "course_id": matched_course.id,
                        "exam_type": exam_type,
                        "module": module_name or None,
                        "semester_name": semester_name or None,
                        "semester_id": semester_id,
                        "phone_number": phone or None,
                        "username": username or admission.replace("/", "").replace(" ", "").lower(),
                        "password": password or admission,
                    })

            session["hod_upload_valid"] = valid_rows

            return render_template(
                "hod/upload_students_preview.html",
                department=department, hod_name=session.get("username"),
                valid_rows=valid_rows, error_rows=error_rows, total_rows=len(df),
            )

        return render_template(
            "hod/upload_students.html",
            department=department, hod_name=session.get("username"),
        )
    finally:
        db.close()


@app.route("/hod/students/upload/confirm", methods=["POST"])
def hod_students_upload_confirm():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    rows = session.pop("hod_upload_valid", None)
    if not rows:
        flash("No validated rows to import. Please upload again.", "warning")
        return redirect(url_for("hod_students_upload"))

    db = SessionLocal()
    imported, skipped = 0, 0
    try:
        course_ids = list({r["course_id"] for r in rows})

        courses = {
            c.id: c for c in db.query(Course).filter(
                Course.id.in_(course_ids),
                Course.department_id == department_id,
            ).all()
        }

        admission_numbers = [r["admission_number"] for r in rows]
        existing_admissions = {
            a[0] for a in db.query(StudentProfile.admission_number)
                           .filter(StudentProfile.admission_number.in_(admission_numbers))
                           .all()
            if a[0]
        }

        base_usernames = [
            (r["username"] or r["admission_number"]).strip()[:45]
            for r in rows
        ]
        existing_usernames = {
            u[0] for u in db.query(User.username)
                           .filter(User.username.in_(base_usernames))
                           .all()
            if u[0]
        }

        current_semesters = {
            s.course_id: s.id
            for s in db.query(Semester).filter(
                Semester.course_id.in_(course_ids),
                Semester.is_current == True,
            ).all()
        }

        modules_lookup = {}
        for m in db.query(Module).filter(Module.course_id.in_(course_ids)).all():
            modules_lookup[(m.course_id, m.name)] = m

        new_users = []
        pending = []
        username_taken = set(existing_usernames)

        for r in rows:
            course = courses.get(r["course_id"])
            if not course:
                skipped += 1
                continue

            adm = r["admission_number"]
            if adm in existing_admissions:
                skipped += 1
                continue

            base_username = (r["username"] or adm).strip()[:45] or adm
            username = base_username
            counter = 1
            while username in username_taken:
                username = f"{base_username}{counter}"
                counter += 1
            username_taken.add(username)

            user = User(username=username, password=r["password"] or adm)

            sem_id = r.get("semester_id")
            if sem_id is None:
                sem_id = current_semesters.get(course.id)

            module_id = None
            module_label = r.get("module")
            if module_label:
                mod = modules_lookup.get((course.id, module_label))
                if mod:
                    module_id = mod.id
                    module_label = mod.name

            new_users.append(user)
            pending.append({
                "course_id": course.id,
                "full_name": r["full_name"],
                "admission_number": adm,
                "exam_type": r["exam_type"],
                "phone_number": r["phone_number"],
                "current_semester_id": sem_id,
                "module": module_label,
                "module_id": module_id,
            })
            existing_admissions.add(adm)

        if new_users:
            db.add_all(new_users)
            db.flush()
            new_students = []
            for user, data in zip(new_users, pending):
                new_students.append(StudentProfile(user_id=user.id, **data))
            db.add_all(new_students)
            db.commit()
            imported = len(new_students)
        else:
            db.commit()

        msg = f"✅ Imported {imported} student(s)."
        if skipped:
            msg += f" Skipped {skipped} (duplicates or course mismatch)."
        flash(msg, "success" if imported else "warning")

    except IntegrityError as e:
        db.rollback()
        import traceback; traceback.print_exc()
        flash(f"❌ Import failed: {_translate_db_error(e)}", "danger")
    except Exception as e:
        db.rollback()
        import traceback
        traceback.print_exc()
        flash(f"❌ Import failed: {e}", "danger")
    finally:
        db.close()

    return redirect(url_for("hod_students_list"))


@app.route("/hod/student/<int:student_id>/set-module", methods=["POST"])
def hod_set_student_module(student_id):
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    semester_id = request.form.get("semester_id", type=int)
    module_value = request.form.get("module", "").strip() or None

    db = SessionLocal()
    try:
        student = db.query(StudentProfile).filter_by(id=student_id).first()
        if not student:
            flash("Student not found.", "danger")
            return redirect(url_for("hod_students_list"))

        course = db.query(Course).filter_by(id=student.course_id).first()
        if not course or course.department_id != department_id:
            flash("Student not in your department.", "danger")
            return redirect(url_for("hod_students_list"))

        if "module" in request.form:
            student.module = module_value

        if "semester_id" in request.form:
            if semester_id:
                sem = db.query(Semester).filter_by(id=semester_id, course_id=student.course_id).first()
                if not sem:
                    flash("That semester doesn't belong to this student's course.", "danger")
                    return redirect(url_for("hod_students_list"))
                student.current_semester_id = sem.id
            else:
                student.current_semester_id = None

        db.commit()
        flash(f"✅ {student.full_name} updated.", "success")
    except Exception as e:
        db.rollback()
        flash(f"Error: {str(e)}", "danger")
    finally:
        db.close()

    return redirect(url_for("hod_students_list"))


@app.route("/hod/students/assign-all-to-current", methods=["POST"])
def hod_assign_all_to_current():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    db = SessionLocal()
    try:
        course_ids = [c.id for c in db.query(Course).filter_by(department_id=department_id).all()]
        current_by_course = {
            s.course_id: s.id
            for s in db.query(Semester).filter(
                Semester.course_id.in_(course_ids),
                Semester.is_current == True,
            ).all()
        }

        updated = 0
        for student in (
            student_query(db)
            .filter(StudentProfile.course_id.in_(course_ids),
                    StudentProfile.current_semester_id.is_(None))
            .all()
        ):
            sem_id = current_by_course.get(student.course_id)
            if sem_id:
                student.current_semester_id = sem_id
                updated += 1

        db.commit()
        flash(f"✅ Assigned {updated} student(s) to their course's current semester.", "success")
    except Exception as e:
        db.rollback()
        flash(f"Error: {e}", "danger")
    finally:
        db.close()

    return redirect(url_for("hod_dashboard"))


@app.route("/hod/student/<int:student_id>/promote", methods=["POST"])
def hod_promote_student(student_id):
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    db = SessionLocal()
    try:
        student = db.query(StudentProfile).filter_by(id=student_id).first()
        if not student:
            flash("Student not found.", "danger")
            return redirect(url_for("hod_students_list"))

        course = db.query(Course).filter_by(id=student.course_id).first()
        if not course or course.department_id != department_id:
            flash("Student not in your department.", "danger")
            return redirect(url_for("hod_students_list"))

        semesters = (
            db.query(Semester).filter_by(course_id=student.course_id)
            .order_by(Semester.created_at.asc()).all()
        )
        if not semesters:
            flash("No semesters defined for this course yet.", "warning")
            return redirect(url_for("hod_students_list"))

        current_sem_id = None
        for sem in semesters:
            has_marks = db.query(KnecMark).filter_by(
                student_id=student.id, semester_id=sem.id,
            ).count() > 0
            if not has_marks:
                current_sem_id = sem.id
                break

        if current_sem_id is None:
            flash(f"{student.full_name} is already in the final semester.", "warning")
            return redirect(url_for("hod_students_list"))

        current_idx = next((i for i, s in enumerate(semesters) if s.id == current_sem_id), None)
        if current_idx is None or current_idx + 1 >= len(semesters):
            flash("Already in the last semester — cannot promote further.", "warning")
            return redirect(url_for("hod_students_list"))

        next_sem = semesters[current_idx + 1]
        student.current_semester_id = next_sem.id
        db.commit()
        flash(f"✅ {student.full_name} promoted to '{next_sem.name}'.", "success")
    except Exception as e:
        db.rollback()
        flash(f"Error promoting student: {str(e)}", "danger")
    finally:
        db.close()

    return redirect(url_for("hod_students_list"))


@app.route("/hod/course/<int:course_id>")
def hod_course_detail(course_id):
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    db = SessionLocal()
    try:
        course = db.query(Course).filter_by(id=course_id, department_id=department_id).first()
        if not course:
            flash("Course not found in your department.", "danger")
            return redirect(url_for("hod_dashboard"))

        students = (
            student_query(db).filter_by(course_id=course_id)
            .order_by(StudentProfile.full_name.asc()).all()
        )

        semesters = (
            semester_query(db).filter_by(course_id=course_id)
            .order_by(Semester.created_at.asc()).all()
        )
        semester_groups = []
        for sem in semesters:
            subs = (
                subject_query(db)
                .filter_by(course_id=course_id, semester_id=sem.id)
                .order_by(Subject.name.asc()).all()
            )
            semester_groups.append({
                "semester": sem, "subjects": subs,
                "is_current": bool(sem.is_current),
            })

        unassigned = (
            subject_query(db)
            .filter(Subject.course_id == course_id)
            .filter(Subject.semester_id.is_(None))
            .order_by(Subject.name.asc()).all()
        )

        active_count = sum(1 for s in students if not s.blocked)
        blocked_count = len(students) - active_count

        return render_template(
            "hod/course_detail.html",
            department=db.query(Department).filter_by(id=department_id).first(),
            course=course, students=students,
            semester_groups=semester_groups, unassigned_subjects=unassigned,
            active_count=active_count, blocked_count=blocked_count,
        )
    finally:
        db.close()


@app.route("/hod/student/<int:student_id>")
def hod_student_detail(student_id):
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    db = SessionLocal()
    try:
        student = student_query(db).filter_by(id=student_id).first()
        if not student:
            flash("Student not found.", "danger")
            return redirect(url_for("hod_students_list"))

        course = db.query(Course).filter_by(id=student.course_id).first()
        if not course or course.department_id != department_id:
            flash("That student is not in your department.", "danger")
            return redirect(url_for("hod_dashboard"))

        atype = get_course_assessment_type(course)

        semesters = (
            semester_query(db).filter_by(course_id=student.course_id)
            .order_by(Semester.created_at.asc()).all()
        )

        semester_reports = []
        grand_total = 0
        grand_count = 0

        for sem in semesters:
            subjects = (
                subject_query(db)
                .filter_by(course_id=student.course_id, semester_id=sem.id)
                .order_by(Subject.name.asc()).all()
            )
            rows = []
            sem_total = 0
            sem_count = 0

            # ── Only include subjects the student actually has marks for ──
            for subj in subjects:
                info = get_student_subject_marks(db, student.id, subj.id, sem.id, course)
                if not info:
                    continue
                rows.append({"subject": subj, "info": info, "has_marks": True})
                sem_total += info["total"]
                sem_count += 1

            avg = round(sem_total / sem_count, 2) if sem_count else None
            semester_reports.append({
                "semester": sem, "rows": rows, "is_current": bool(sem.is_current),
                "sem_total": sem_total, "sem_count": sem_count, "avg": avg,
            })
            grand_total += sem_total
            grand_count += sem_count

        overall_avg = round(grand_total / grand_count, 2) if grand_count else None
        labels = list(KASNEB_LABELS) if atype == "kasneb" else get_periodic_labels(course)

        return render_template(
            "hod/student_detail.html",
            department=db.query(Department).filter_by(id=department_id).first(),
            student=student, course=course, semester_reports=semester_reports,
            grand_total=grand_total, grand_count=grand_count, overall_avg=overall_avg,
            assessment_type=atype,
            periodic_labels=labels,
        )
    finally:
        db.close()


@app.route("/hod/all-students", methods=["GET"])
def hod_all_students():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    db = SessionLocal()
    try:
        department = db.query(Department).filter_by(id=department_id).first()
        if not department:
            flash("Department not found.", "danger")
            return redirect(url_for("logout"))

        courses = db.query(Course).filter_by(department_id=department_id).order_by(Course.name.asc()).all()
        course_ids = [c.id for c in courses]

        search = request.args.get("search", "").strip()
        course_id = request.args.get("course_id", type=int)
        level = request.args.get("level", "").strip()
        exam_type = request.args.get("exam_type", "").strip()

        query = student_query(db)
        if course_ids:
            query = query.filter(StudentProfile.course_id.in_(course_ids))
        else:
            query = query.filter(False)

        if search:
            like = f"%{search}%"
            query = query.filter(or_(
                StudentProfile.full_name.ilike(like),
                StudentProfile.admission_number.ilike(like),
                StudentProfile.phone_number.ilike(like),
            ))
        if course_id and course_id in course_ids:
            query = query.filter(StudentProfile.course_id == course_id)
        if level:
            query = query.join(Course, StudentProfile.course_id == Course.id) \
                         .filter(Course.level.ilike(f"%{level}%"))
        if exam_type:
            query = query.filter(StudentProfile.exam_type == exam_type)

        students = query.order_by(StudentProfile.full_name.asc()).all()

        levels_present = sorted({
            c.level.strip() for c in courses if c.level and c.level.strip()
        })

        semester_map = {}
        for cid in course_ids:
            sems = db.query(Semester).filter_by(course_id=cid).order_by(Semester.created_at.asc()).all()
            semester_map[cid] = [{"id": s.id, "name": s.name} for s in sems]

        return render_template(
            "hod/all_students.html",
            department=department, students=students, courses=courses,
            levels_present=levels_present, semester_map=semester_map,
            hod_name=session.get("username"),
            search=search, course_id=course_id, level=level, exam_type=exam_type,
        )
    finally:
        db.close()


@app.route("/hod/all-courses", methods=["GET"])
def hod_all_courses():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    db = SessionLocal()
    try:
        department = db.query(Department).filter_by(id=department_id).first()
        if not department:
            flash("Department not found.", "danger")
            return redirect(url_for("logout"))

        courses = db.query(Course).filter_by(department_id=department_id).order_by(Course.name.asc()).all()

        courses_data = []
        for course in courses:
            students_count = db.query(StudentProfile).filter_by(course_id=course.id).count()
            semesters = (
                semester_query(db).filter_by(course_id=course.id)
                .order_by(Semester.created_at.asc()).all()
            )

            semesters_data = []
            total_subjects = 0
            for sem in semesters:
                subs = (
                    subject_query(db).filter_by(course_id=course.id, semester_id=sem.id)
                    .order_by(Subject.name.asc()).all()
                )
                total_subjects += len(subs)
                semesters_data.append({"semester": sem, "subjects": subs})

            unassigned = (
                subject_query(db).filter(Subject.course_id == course.id)
                .filter(Subject.semester_id.is_(None))
                .order_by(Subject.name.asc()).all()
            )
            total_subjects += len(unassigned)

            courses_data.append({
                "course": course, "students_count": students_count,
                "semesters": semesters_data, "unassigned": unassigned,
                "total_subjects": total_subjects,
            })

        return render_template(
            "hod/all_courses.html",
            department=department, courses=courses_data,
            hod_name=session.get("username"),
        )
    finally:
        db.close()


@app.route("/hod/courses", methods=["GET"])
def hod_courses_search():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    db = SessionLocal()
    try:
        department = db.query(Department).filter_by(id=department_id).first()
        if not department:
            flash("Department not found.", "danger")
            return redirect(url_for("logout"))

        search_query_str = request.args.get("search", "").strip()
        q = db.query(Course).filter_by(department_id=department_id)
        if search_query_str:
            like = f"%{search_query_str}%"
            q = q.filter(or_(Course.name.ilike(like), Course.level.ilike(like)))

        courses = q.order_by(Course.name.asc()).all()

        course_rows = []
        for c in courses:
            students_count = db.query(StudentProfile).filter_by(course_id=c.id).count()
            subjects_count = db.query(Subject).filter_by(course_id=c.id).count()
            semesters_count = db.query(Semester).filter_by(course_id=c.id).count()
            current_sem = db.query(Semester).filter_by(course_id=c.id, is_current=True).first()

            course_rows.append({
                "course": c, "students_count": students_count,
                "subjects_count": subjects_count, "semesters_count": semesters_count,
                "current_semester": current_sem,
            })

        return render_template(
            "hod/courses_list.html",
            department=department, hod_name=session.get("username"),
            course_rows=course_rows, search_query=search_query_str,
        )
    finally:
        db.close()


@app.route("/hod/marks")
def hod_marks_picker():
    if session.get("role") != "hod":
        return redirect(url_for("login"))
    return redirect(url_for("hod_marks_department", dept_id=session["department_id"]))


@app.route("/hod/marks/department/<int:dept_id>", methods=["GET"])
def hod_marks_department(dept_id):
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    if dept_id != session["department_id"]:
        flash("You can only manage your own department.", "danger")
        return redirect(url_for("hod_dashboard"))

    db = SessionLocal()
    try:
        department = db.query(Department).filter_by(id=dept_id).first()
        if not department:
            flash("Department not found.", "danger")
            return redirect(url_for("logout"))

        courses = db.query(Course).filter_by(department_id=dept_id).order_by(Course.name.asc()).all()

        course_id = request.args.get("course_id", type=int)
        semester_id = request.args.get("semester_id", type=int)
        module_id = request.args.get("module_id", type=int)

        course = semester = None
        subjects, semesters, available_modules = [], [], []

        if course_id:
            course = db.query(Course).filter_by(id=course_id, department_id=dept_id).first()
        if course:
            semesters = (
                semester_query(db).filter_by(course_id=course.id)
                .order_by(Semester.created_at.desc()).all()
            )
            if semester_id:
                semester = next((s for s in semesters if s.id == semester_id), None)
            elif semesters:
                semester = next((s for s in semesters if s.is_current), semesters[0])

            if semester:
                all_sem_subjects = (
                    subject_query(db)
                    .filter_by(course_id=course.id, semester_id=semester.id)
                    .order_by(Subject.name.asc()).all()
                )
                used_module_ids = {s.module_id for s in all_sem_subjects if s.module_id}
                if used_module_ids:
                    available_modules = (
                        db.query(Module).filter(Module.id.in_(used_module_ids))
                        .order_by(Module.order.asc(), Module.name.asc()).all()
                    )
                subjects = (
                    [s for s in all_sem_subjects if s.module_id == module_id]
                    if module_id else all_sem_subjects
                )

        return render_template(
            "admin/marks_step2_course.html",
            department=department, courses=courses,
            course=course, semesters=semesters, semester=semester,
            subjects=subjects, available_modules=available_modules,
            module_id=module_id, hod_mode=True,
        )
    finally:
        db.close()


@app.route("/hod/marks/enter", methods=["GET", "POST"])
def hod_enter_marks():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        department_id = session["department_id"]
        src = request.form if request.method == "POST" else request.args
        subject_id = src.get("subject_id", type=int)
        semester_id = src.get("semester_id", type=int)

        if not subject_id:
            flash("Pick a subject first.", "danger")
            return redirect(url_for("hod_marks_picker"))

        subject = subject_query(db).filter_by(id=subject_id).first()
        if not subject:
            flash("Subject not found.", "danger")
            return redirect(url_for("hod_marks_picker"))

        course = subject.course
        if not course or course.department_id != department_id:
            flash("You can only manage marks for your department.", "danger")
            return redirect(url_for("hod_dashboard"))

        semester = None
        if semester_id:
            semester = db.query(Semester).filter_by(id=semester_id).first()
            if not semester:
                flash("Semester not found.", "danger")
                return redirect(url_for("hod_marks_picker"))
            if subject.course_id != semester.course_id:
                flash("Subject and semester belong to different courses.", "danger")
                return redirect(url_for("hod_marks_picker"))

        atype = get_course_assessment_type(course)

        all_semesters = (
            db.query(Semester)
            .filter_by(course_id=course.id)
            .order_by(Semester.created_at.desc())
            .all()
        )

        students_q = student_query(db).filter_by(course_id=subject.course_id)
        fresh_module_id = db.query(Subject.module_id).filter(Subject.id == subject_id).scalar()

        if fresh_module_id:
            students_q = students_q.filter(
                or_(
                    StudentProfile.module_id == fresh_module_id,
                    StudentProfile.module_id.is_(None),
                )
            )

        students = students_q.order_by(StudentProfile.full_name.asc()).all()

        if request.method == "POST":
            def to_float(v):
                v = (v or "").strip()
                if not v:
                    return None
                try:
                    return float(v)
                except ValueError:
                    return None

            def to_int(v):
                v = (v or "").strip()
                return int(v) if v.isdigit() else None

            saved = 0

            if atype == "short":
                max_s = course.max_score or 100
                for s in students:
                    sc = to_float(request.form.get(f"score_{s.id}"))
                    rm = (request.form.get(f"remarks_{s.id}") or "").strip() or None
                    if sc is not None and not (0 <= sc <= max_s):
                        flash(f"Score for {s.full_name} must be 0–{max_s}.", "warning")
                        continue
                    existing = db.query(ShortCourseMark).filter_by(
                        student_id=s.id, subject_id=subject_id, semester_id=semester_id
                    ).first()
                    if sc is None and not rm:
                        if existing:
                            db.delete(existing)
                        continue
                    if existing:
                        existing.score = sc
                        existing.remarks = rm
                        existing.entered_at = datetime.utcnow()
                    else:
                        db.add(ShortCourseMark(
                            student_id=s.id, subject_id=subject_id,
                            semester_id=semester_id, score=sc, remarks=rm,
                        ))
                    saved += 1

            elif atype == "kasneb":
                max_s = course.max_score or 100
                for s in students:
                    c1 = to_float(request.form.get(f"score_{s.id}__CAT 1"))
                    c2 = to_float(request.form.get(f"score_{s.id}__CAT 2"))
                    if c1 is not None and not (0 <= c1 <= max_s):
                        flash(f"CAT 1 for {s.full_name} must be 0–{max_s}.", "warning")
                        continue
                    if c2 is not None and not (0 <= c2 <= max_s):
                        flash(f"CAT 2 for {s.full_name} must be 0–{max_s}.", "warning")
                        continue
                    _save_kasneb_marks(db, s.id, subject_id, semester_id, c1, c2)
                    if c1 is not None or c2 is not None:
                        saved += 1

            elif atype == "periodic":
                max_s = course.max_score or 100
                labels = get_periodic_labels(course)
                for s in students:
                    for label in labels:
                        sc = to_float(request.form.get(f"score_{s.id}__{label}"))
                        rm = (request.form.get(f"remarks_{s.id}__{label}") or "").strip() or None
                        if sc is not None and not (0 <= sc <= max_s):
                            flash(f"'{label}' for {s.full_name} must be 0–{max_s}.", "warning")
                            continue
                        existing = db.query(PeriodicMark).filter_by(
                            student_id=s.id, subject_id=subject_id,
                            semester_id=semester_id, label=label,
                        ).first()
                        if sc is None and not rm:
                            if existing:
                                db.delete(existing)
                            continue
                        if existing:
                            existing.score = sc
                            existing.remarks = rm
                            existing.entered_at = datetime.utcnow()
                        else:
                            db.add(PeriodicMark(
                                student_id=s.id, subject_id=subject_id,
                                semester_id=semester_id, label=label,
                                score=sc, remarks=rm,
                            ))
                        saved += 1

            else:  # KNEC
                for s in students:
                    c1 = to_int(request.form.get(f"cat1_{s.id}"))
                    c2 = to_int(request.form.get(f"cat2_{s.id}"))
                    fn = to_int(request.form.get(f"final_{s.id}"))
                    existing = db.query(KnecMark).filter_by(
                        student_id=s.id, subject_id=subject_id, semester_id=semester_id
                    ).first()
                    if c1 is None and c2 is None and fn is None:
                        if existing:
                            db.delete(existing)
                        continue
                    if c1 is not None and not (0 <= c1 <= 15):
                        flash(f"CAT1 for {s.full_name} must be 0–15.", "warning")
                        continue
                    if c2 is not None and not (0 <= c2 <= 15):
                        flash(f"CAT2 for {s.full_name} must be 0–15.", "warning")
                        continue
                    if fn is not None and not (0 <= fn <= 70):
                        flash(f"Final for {s.full_name} must be 0–70.", "warning")
                        continue
                    if existing:
                        existing.cat1, existing.cat2, existing.final = c1, c2, fn
                        existing.entered_at = datetime.utcnow()
                    else:
                        db.add(KnecMark(
                            student_id=s.id, subject_id=subject_id,
                            semester_id=semester_id,
                            cat1=c1, cat2=c2, final=fn,
                        ))
                    saved += 1

            db.commit()
            flash(f"✅ Saved marks for {saved} record(s).", "success")
            return redirect(url_for(
                "hod_marks_department",
                dept_id=department_id,
                course_id=subject.course_id,
                semester_id=semester_id,
            ))

        if atype == "short":
            existing = {m.student_id: m for m in db.query(ShortCourseMark).filter_by(
                subject_id=subject_id, semester_id=semester_id).all()}
        elif atype in ("kasneb", "periodic"):
            existing = {}
            for m in db.query(PeriodicMark).filter_by(
                subject_id=subject_id, semester_id=semester_id).all():
                existing.setdefault(m.student_id, {})[m.label] = m
        else:
            existing = {m.student_id: m for m in db.query(KnecMark).filter_by(
                subject_id=subject_id, semester_id=semester_id).all()}

        labels = list(KASNEB_LABELS) if atype == "kasneb" else get_periodic_labels(course)

        return render_template(
            "admin/marks_entry.html",
            subject=subject, semester=semester, course=course,
            department=course.department, students=students, existing=existing,
            assessment_type=atype,
            periodic_labels=labels,
            hod_mode=True,
            all_semesters=all_semesters,
        )
    finally:
        db.close()


@app.route("/hod/marks/class")
def hod_class_marks():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    semester_id = request.args.get("semester_id", type=int)

    db = SessionLocal()
    try:
        if not semester_id:
            flash("Pick a semester.", "warning")
            return redirect(url_for("hod_marks_picker"))

        semester = semester_query(db).filter_by(id=semester_id).first()
        if not semester:
            flash("Semester not found.", "danger")
            return redirect(url_for("hod_marks_picker"))

        course = semester.course
        if not course or course.department_id != department_id:
            flash("That semester is not in your department.", "danger")
            return redirect(url_for("hod_dashboard"))

        atype = get_course_assessment_type(course)

        students = (
            student_query(db).filter_by(course_id=semester.course_id)
            .order_by(StudentProfile.full_name.asc()).all()
        )
        subjects = (
            subject_query(db).filter_by(course_id=semester.course_id, semester_id=semester_id)
            .order_by(Subject.name.asc()).all()
        )

        if atype == "short":
            marks = {}
            for m in db.query(ShortCourseMark).filter_by(semester_id=semester_id).all():
                marks.setdefault(m.student_id, {})[m.subject_id] = m
        elif atype in ("kasneb", "periodic"):
            marks = {}
            for m in db.query(PeriodicMark).filter_by(semester_id=semester_id).all():
                marks.setdefault(m.student_id, {}).setdefault(m.subject_id, {})[m.label] = m
        else:
            marks = {}
            for m in db.query(KnecMark).filter_by(semester_id=semester_id).all():
                marks.setdefault(m.student_id, {})[m.subject_id] = m

        labels = list(KASNEB_LABELS) if atype == "kasneb" else get_periodic_labels(course)

        return render_template(
            "admin/class_marks.html",
            semester=semester, students=students, subjects=subjects, marks=marks,
            assessment_type=atype,
            periodic_labels=labels,
            hod_mode=True,
        )
    finally:
        db.close()


@app.route("/hod/student/<int:student_id>/marks", methods=["GET", "POST"])
def hod_student_marks(student_id):
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    db = SessionLocal()
    try:
        student = db.query(StudentProfile).filter_by(id=student_id).first()
        if not student:
            flash("Student not found.", "danger")
            return redirect(url_for("hod_students_list"))

        course = db.query(Course).filter_by(id=student.course_id).first()
        if not course or course.department_id != department_id:
            flash("That student is not in your department.", "danger")
            return redirect(url_for("hod_dashboard"))

        atype = get_course_assessment_type(course)

        semester_id = (
            request.args.get("semester_id", type=int)
            or request.form.get("semester_id", type=int)
        )

        semesters = (
            semester_query(db).filter_by(course_id=student.course_id)
            .order_by(Semester.created_at.desc()).all()
        )

        semester = None
        if semester_id:
            semester = next((s for s in semesters if s.id == semester_id), None)
        elif semesters:
            semester = next((s for s in semesters if s.is_current), semesters[0])

        subjects = []
        if semester:
            subjects = (
                subject_query(db).filter_by(course_id=student.course_id, semester_id=semester.id)
                .order_by(Subject.name.asc()).all()
            )

        if request.method == "POST" and semester:
            def to_float(v):
                v = (v or "").strip()
                if not v:
                    return None
                try:
                    return float(v)
                except ValueError:
                    return None

            def to_int(v):
                v = (v or "").strip()
                return int(v) if v.isdigit() else None

            saved = 0

            if atype == "short":
                max_s = course.max_score or 100
                for subj in subjects:
                    sc = to_float(request.form.get(f"score_{subj.id}"))
                    rm = (request.form.get(f"remarks_{subj.id}") or "").strip() or None

                    if sc is not None and not (0 <= sc <= max_s):
                        flash(f"Score for {subj.name} must be 0–{max_s}.", "warning")
                        continue

                    existing = db.query(ShortCourseMark).filter_by(
                        student_id=student.id, subject_id=subj.id, semester_id=semester.id
                    ).first()

                    if sc is None and not rm:
                        if existing:
                            db.delete(existing)
                        continue

                    if existing:
                        existing.score = sc
                        existing.remarks = rm
                        existing.entered_at = datetime.utcnow()
                    else:
                        db.add(ShortCourseMark(
                            student_id=student.id, subject_id=subj.id,
                            semester_id=semester.id, score=sc, remarks=rm,
                        ))
                    saved += 1

            elif atype == "kasneb":
                max_s = course.max_score or 100
                for subj in subjects:
                    c1 = to_float(request.form.get(f"score_{subj.id}__CAT 1"))
                    c2 = to_float(request.form.get(f"score_{subj.id}__CAT 2"))

                    if c1 is not None and not (0 <= c1 <= max_s):
                        flash(f"CAT 1 for {subj.name} must be 0–{max_s}.", "warning")
                        continue
                    if c2 is not None and not (0 <= c2 <= max_s):
                        flash(f"CAT 2 for {subj.name} must be 0–{max_s}.", "warning")
                        continue

                    _save_kasneb_marks(db, student.id, subj.id, semester.id, c1, c2)
                    if c1 is not None or c2 is not None:
                        saved += 1

            elif atype == "periodic":
                max_s = course.max_score or 100
                labels = get_periodic_labels(course)
                for subj in subjects:
                    for label in labels:
                        sc = to_float(request.form.get(f"score_{subj.id}__{label}"))
                        rm = (request.form.get(f"remarks_{subj.id}__{label}") or "").strip() or None

                        if sc is not None and not (0 <= sc <= max_s):
                            flash(f"'{label}' for {subj.name} must be 0–{max_s}.", "warning")
                            continue

                        existing = db.query(PeriodicMark).filter_by(
                            student_id=student.id, subject_id=subj.id,
                            semester_id=semester.id, label=label,
                        ).first()

                        if sc is None and not rm:
                            if existing:
                                db.delete(existing)
                            continue

                        if existing:
                            existing.score = sc
                            existing.remarks = rm
                            existing.entered_at = datetime.utcnow()
                        else:
                            db.add(PeriodicMark(
                                student_id=student.id, subject_id=subj.id,
                                semester_id=semester.id, label=label,
                                score=sc, remarks=rm,
                            ))
                        saved += 1

            else:  # KNEC
                for subj in subjects:
                    c1 = to_int(request.form.get(f"cat1_{subj.id}"))
                    c2 = to_int(request.form.get(f"cat2_{subj.id}"))
                    fn = to_int(request.form.get(f"final_{subj.id}"))

                    existing = db.query(KnecMark).filter_by(
                        student_id=student.id, subject_id=subj.id, semester_id=semester.id
                    ).first()

                    if c1 is None and c2 is None and fn is None:
                        if existing:
                            db.delete(existing)
                        continue

                    if existing:
                        existing.cat1, existing.cat2, existing.final = c1, c2, fn
                        existing.entered_at = datetime.utcnow()
                    else:
                        db.add(KnecMark(
                            student_id=student.id, subject_id=subj.id, semester_id=semester.id,
                            cat1=c1, cat2=c2, final=fn,
                        ))
                    saved += 1

            db.commit()
            flash(f"✅ Saved marks for {saved} record(s).", "success")
            return redirect(url_for(
                "hod_student_marks",
                student_id=student.id, semester_id=semester.id,
            ))

        rows = []

        if atype == "short":
            marks = {
                m.subject_id: m
                for m in db.query(ShortCourseMark).filter_by(
                    student_id=student.id,
                    semester_id=semester.id if semester else None,
                ).all()
            }
            for subj in subjects:
                m = marks.get(subj.id)
                rows.append({
                    "subject": subj,
                    "score": m.score if m else None,
                    "remarks": m.remarks if m else None,
                })

        elif atype in ("kasneb", "periodic"):
            marks = {}
            for m in db.query(PeriodicMark).filter_by(
                student_id=student.id,
                semester_id=semester.id if semester else None,
            ).all():
                marks.setdefault(m.subject_id, {})[m.label] = m.score

            for subj in subjects:
                rows.append({
                    "subject": subj,
                    "cells": marks.get(subj.id, {}),
                })

        else:  # KNEC
            marks = {
                m.subject_id: m
                for m in db.query(KnecMark).filter_by(
                    student_id=student.id,
                    semester_id=semester.id if semester else None,
                ).all()
            }
            for subj in subjects:
                m = marks.get(subj.id)
                rows.append({
                    "subject": subj,
                    "cat1": m.cat1 if m and m.cat1 is not None else None,
                    "cat2": m.cat2 if m and m.cat2 is not None else None,
                    "final": m.final if m and m.final is not None else None,
                })

        labels = list(KASNEB_LABELS) if atype == "kasneb" else get_periodic_labels(course)

        return render_template(
            "hod/student_marks.html",
            student=student, course=course, semester=semester,
            semesters=semesters, rows=rows,
            assessment_type=atype,
            periodic_labels=labels,
        )
    finally:
        db.close()


@app.route("/hod/marks/class/export")
def hod_class_marks_export():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    semester_id = request.args.get("semester_id", type=int)
    if not semester_id:
        flash("Pick a semester.", "warning")
        return redirect(url_for("hod_marks_picker"))

    db = SessionLocal()
    try:
        semester = semester_query(db).filter_by(id=semester_id).first()
        if not semester:
            flash("Semester not found.", "danger")
            return redirect(url_for("hod_dashboard"))

        course = semester.course
        if not course or course.department_id != department_id:
            flash("That semester is not in your department.", "danger")
            return redirect(url_for("hod_dashboard"))

        atype = get_course_assessment_type(course)

        students = (
            student_query(db).filter_by(course_id=semester.course_id)
            .order_by(StudentProfile.full_name.asc()).all()
        )
        subjects = (
            subject_query(db).filter_by(course_id=semester.course_id, semester_id=semester_id)
            .order_by(Subject.name.asc()).all()
        )

        data = []
        if atype == "knec":
            marks = {}
            for m in db.query(KnecMark).filter_by(semester_id=semester_id).all():
                marks.setdefault(m.student_id, {})[m.subject_id] = m
            for s in students:
                row = {"Admission No": s.admission_number, "Student Name": s.full_name}
                grand_total = counted = 0
                for subj in subjects:
                    m = marks.get(s.id, {}).get(subj.id)
                    c1 = m.cat1 if m and m.cat1 is not None else ""
                    c2 = m.cat2 if m and m.cat2 is not None else ""
                    fn = m.final if m and m.final is not None else ""
                    total = (m.cat1 or 0) + (m.cat2 or 0) + (m.final or 0) if m else ""
                    prefix = subj.name.split(" ", 1)[0]
                    row[f"{prefix} CAT1"] = c1
                    row[f"{prefix} CAT2"] = c2
                    row[f"{prefix} Final"] = fn
                    row[f"{prefix} Total"] = total
                    if isinstance(total, int):
                        grand_total += total
                        counted += 1
                row["Grand Total"] = grand_total
                row["Average"] = round(grand_total / counted, 2) if counted else ""
                data.append(row)

        elif atype == "short":
            marks = {}
            for m in db.query(ShortCourseMark).filter_by(semester_id=semester_id).all():
                marks.setdefault(m.student_id, {})[m.subject_id] = m
            for s in students:
                row = {"Admission No": s.admission_number, "Student Name": s.full_name}
                grand_total = counted = 0
                for subj in subjects:
                    m = marks.get(s.id, {}).get(subj.id)
                    score = m.score if m and m.score is not None else ""
                    row[f"{subj.name} Score"] = score
                    row[f"{subj.name} Remarks"] = m.remarks if m and m.remarks else ""
                    if isinstance(score, (int, float)):
                        grand_total += score
                        counted += 1
                row["Grand Total"] = grand_total
                row["Average"] = round(grand_total / counted, 2) if counted else ""
                data.append(row)

        elif atype == "kasneb":
            marks = {}
            for m in db.query(PeriodicMark).filter_by(semester_id=semester_id).all():
                marks.setdefault(m.student_id, {}).setdefault(m.subject_id, {})[m.label] = m
            for s in students:
                row = {"Admission No": s.admission_number, "Student Name": s.full_name}
                for subj in subjects:
                    cells = marks.get(s.id, {}).get(subj.id, {})
                    row[f"{subj.name} CAT1"] = cells["CAT 1"].score if "CAT 1" in cells else ""
                    row[f"{subj.name} CAT2"] = cells["CAT 2"].score if "CAT 2" in cells else ""
                data.append(row)

        else:  # periodic
            labels = get_periodic_labels(course)
            marks = {}
            for m in db.query(PeriodicMark).filter_by(semester_id=semester_id).all():
                marks.setdefault(m.student_id, {}).setdefault(m.subject_id, {})[m.label] = m
            for s in students:
                row = {"Admission No": s.admission_number, "Student Name": s.full_name}
                for subj in subjects:
                    cells = marks.get(s.id, {}).get(subj.id, {})
                    for label in labels:
                        row[f"{subj.name} {label}"] = cells[label].score if label in cells else ""
                data.append(row)

        df = pd.DataFrame(data)
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Marks")
        buf.seek(0)

        safe_course = (course.name if course else "Class").replace(" ", "_")[:30]
        safe_sem = semester.name.replace(" ", "_")

        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"Marks_{safe_course}_{safe_sem}.xlsx",
        )
    finally:
        db.close()


@app.route("/hod/reports/class", methods=["GET"])
def hod_reports_class():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    semester_id = request.args.get("semester_id", type=int)
    if not semester_id:
        flash("Pick a semester first.", "warning")
        return redirect(url_for("hod_marks_picker"))

    db = SessionLocal()
    try:
        semester = semester_query(db).filter_by(id=semester_id).first()
        if not semester:
            flash("Semester not found.", "danger")
            return redirect(url_for("hod_dashboard"))

        course = semester.course
        if not course or course.department_id != department_id:
            flash("That semester is not in your department.", "danger")
            return redirect(url_for("hod_dashboard"))

        atype = get_course_assessment_type(course)

        students = (
            student_query(db).filter_by(course_id=semester.course_id)
            .order_by(StudentProfile.full_name.asc()).all()
        )
        subjects = (
            subject_query(db).filter_by(course_id=semester.course_id, semester_id=semester_id)
            .order_by(Subject.name.asc()).all()
        )

        rows = []
        for s in students:
            row = {
                "student": s, "cells": [],
                "grand_total": 0, "counted": 0,
                "average": None, "grade": None, "result": None,
            }
            for subj in subjects:
                info = get_student_subject_marks(db, s.id, subj.id, semester_id, course)
                row["cells"].append({"subject": subj, "info": info})
                if info:
                    row["grand_total"] += info["total"]
                    row["counted"] += 1

            if row["counted"]:
                row["average"] = round(row["grand_total"] / row["counted"], 2)
                avg = row["average"]
                if atype == "knec":
                    row["grade"] = (1 if avg >= 84 else 2 if avg >= 75 else 3 if avg >= 66
                                    else 4 if avg >= 60 else 5 if avg >= 50
                                    else 6 if avg >= 40 else 7)
                    row["result"] = ("DISTINCTION" if avg >= 75 else "CREDIT" if avg >= 60
                                     else "PASS" if avg >= 40 else "REFER")
                else:
                    row["grade"] = (1 if avg >= 80 else 2 if avg >= 65 else 3 if avg >= 50 else 4)
                    row["result"] = ("DISTINCTION" if avg >= 80 else "CREDIT" if avg >= 65
                                     else "PASS" if avg >= 50 else "FAIL")
            rows.append(row)

        labels = list(KASNEB_LABELS) if atype == "kasneb" else get_periodic_labels(course)

        return render_template(
            "hod/reports_class.html",
            department=db.query(Department).filter_by(id=department_id).first(),
            course=course, semester=semester, subjects=subjects, rows=rows,
            assessment_type=atype,
            periodic_labels=labels,
            hod_name=session.get("username"),
        )
    finally:
        db.close()


@app.route("/hod/reports/class/export")
def hod_reports_class_export():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    semester_id = request.args.get("semester_id", type=int)
    if not semester_id:
        flash("Pick a semester.", "warning")
        return redirect(url_for("hod_marks_picker"))

    db = SessionLocal()
    try:
        semester = semester_query(db).filter_by(id=semester_id).first()
        if not semester:
            flash("Semester not found.", "danger")
            return redirect(url_for("hod_dashboard"))

        course = semester.course
        if not course or course.department_id != department_id:
            flash("That semester is not in your department.", "danger")
            return redirect(url_for("hod_dashboard"))

        atype = get_course_assessment_type(course)

        students = (
            student_query(db).filter_by(course_id=semester.course_id)
            .order_by(StudentProfile.full_name.asc()).all()
        )

        data = []
        for s in students:
            total_sum = 0
            subjects_count = 0

            if atype == "knec":
                for m in db.query(KnecMark).filter_by(student_id=s.id, semester_id=semester_id).all():
                    total_sum += (m.cat1 or 0) + (m.cat2 or 0) + (m.final or 0)
                    subjects_count += 1
            elif atype == "short":
                for m in db.query(ShortCourseMark).filter_by(student_id=s.id, semester_id=semester_id).all():
                    if m.score is not None:
                        total_sum += m.score
                        subjects_count += 1
            else:
                rows = db.query(PeriodicMark).filter_by(student_id=s.id, semester_id=semester_id).all()
                by_subject = {}
                for r in rows:
                    by_subject.setdefault(r.subject_id, []).append(r.score)
                for scores in by_subject.values():
                    valid = [sc for sc in scores if sc is not None]
                    if valid:
                        total_sum += sum(valid) / len(valid)
                        subjects_count += 1

            avg = round(total_sum / subjects_count, 2) if subjects_count else ""

            if not isinstance(avg, float):
                grade = result = ""
            elif atype == "knec":
                grade = (1 if avg >= 84 else 2 if avg >= 75 else 3 if avg >= 66
                         else 4 if avg >= 60 else 5 if avg >= 50
                         else 6 if avg >= 40 else 7)
                result = ("DISTINCTION" if avg >= 75 else "CREDIT" if avg >= 60
                          else "PASS" if avg >= 40 else "REFER")
            else:
                grade = (1 if avg >= 80 else 2 if avg >= 65 else 3 if avg >= 50 else 4)
                result = ("DISTINCTION" if avg >= 80 else "CREDIT" if avg >= 65
                          else "PASS" if avg >= 50 else "FAIL")

            data.append({
                "Admission No": s.admission_number,
                "Student Name": s.full_name,
                "Course": course.name if course else "",
                "Semester": semester.name,
                "Exam Type": (s.exam_type or "").upper(),
                "Assessment Type": atype.upper(),
                "Subjects Graded": subjects_count,
                "Total Marks": total_sum,
                "Average": avg, "Grade": grade, "Overall Result": result,
            })

        df = pd.DataFrame(data)
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Class Report")
        buf.seek(0)

        safe_course = (course.name if course else "Class").replace(" ", "_")[:30]
        safe_sem = semester.name.replace(" ", "_")

        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"ClassReport_{safe_course}_{safe_sem}.xlsx",
        )
    finally:
        db.close()


@app.route("/hod/analytics/knec/<int:course_id>/export")
def hod_knec_analytics_export(course_id):
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    db = SessionLocal()
    try:
        course = db.query(Course).filter_by(
            id=course_id, department_id=department_id
        ).first()
        if not course or get_course_assessment_type(course) != "knec":
            flash("KNEC course not found in your department.", "danger")
            return redirect(url_for("hod_dashboard"))

        buf = _build_knec_analytics_workbook(db, course)

        safe_course = course.name.replace(" ", "_")[:30]
        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"KNEC_Analytics_{safe_course}.xlsx",
        )
    finally:
        db.close()


@app.route("/hod/analytics/knec", methods=["GET"])
def hod_knec_analytics_picker():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    db = SessionLocal()
    try:
        courses = (
            db.query(Course)
            .filter_by(department_id=department_id)
            .order_by(Course.name.asc())
            .all()
        )
        knec_courses = [
            c for c in courses
            if (c.assessment_type or "knec").strip().lower() == "knec"
        ]

        course_rows = []
        for c in knec_courses:
            students_count = db.query(StudentProfile).filter_by(course_id=c.id).count()
            semesters_count = db.query(Semester).filter_by(course_id=c.id).count()
            subjects_count = db.query(Subject).filter_by(course_id=c.id).count()
            marks_count = (
                db.query(KnecMark)
                .join(StudentProfile, KnecMark.student_id == StudentProfile.id)
                .filter(StudentProfile.course_id == c.id)
                .count()
            )
            course_rows.append({
                "course": c,
                "students_count": students_count,
                "semesters_count": semesters_count,
                "subjects_count": subjects_count,
                "marks_count": marks_count,
            })

        return render_template(
            "hod/knec_analytics_picker.html",
            department=db.query(Department).filter_by(id=department_id).first(),
            course_rows=course_rows,
            hod_name=session.get("username"),
        )
    finally:
        db.close()


# ═════════════════════════════════════════════════════════════
# HOD FILTERED REPORT
# ═════════════════════════════════════════════════════════════
@app.route("/hod/reports/filtered_export", methods=["GET", "POST"])
def hod_filtered_report_export():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    db = SessionLocal()
    try:
        course_id   = request.args.get("course_id", type=int)
        semester_id = request.args.get("semester_id", type=int)
        module_id   = request.args.get("module_id", type=int)
        export      = request.args.get("export")

        courses = (
            db.query(Course)
            .filter_by(department_id=department_id)
            .order_by(Course.name.asc())
            .all()
        )
        course_ids = [c.id for c in courses]

        semesters = (
            semester_query(db)
            .filter(Semester.course_id.in_(course_ids))
            .order_by(Semester.course_id.asc(), Semester.created_at.asc())
            .all()
        ) if course_ids else []

        modules = (
            db.query(Module)
            .filter(Module.course_id.in_(course_ids))
            .order_by(Module.course_id.asc(), Module.order.asc(), Module.name.asc())
            .all()
        ) if course_ids else []

        selected_course   = None
        selected_semester = None
        selected_module   = None
        report_rows       = []

        if course_id:
            selected_course = db.query(Course).filter_by(
                id=course_id, department_id=department_id
            ).first()
        if semester_id and selected_course:
            selected_semester = semester_query(db).filter_by(
                id=semester_id, course_id=selected_course.id
            ).first()
        if module_id and selected_course:
            selected_module = db.query(Module).filter_by(
                id=module_id, course_id=selected_course.id
            ).first()

        if selected_course and selected_semester:
            report_rows = _build_filtered_report_rows(
                db, selected_course, selected_semester, selected_module
            )

        if export == "true":
            if not report_rows:
                flash("No data to export for that selection.", "warning")
                return redirect(url_for("hod_filtered_report_export",
                                        course_id=course_id,
                                        semester_id=semester_id,
                                        module_id=module_id))
            return _send_filtered_report_excel(
                report_rows, selected_course, selected_semester, selected_module,
                header_color="168B83", prefix="HOD_Report"
            )

        return render_template(
            "hod/filtered_report.html",
            department=db.query(Department).filter_by(id=department_id).first(),
            hod_name=session.get("username"),
            courses=courses,
            semesters=semesters,
            modules=modules,
            selected_course=selected_course,
            selected_semester=selected_semester,
            selected_module=selected_module,
            report_rows=report_rows,
        )
    finally:
        db.close()


@app.route("/hod/reports/student/<int:student_id>/<int:semester_id>")
def hod_report_student(student_id, semester_id):
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    db = SessionLocal()
    try:
        student = student_query(db).filter_by(id=student_id).first()
        if not student:
            return "Student not found", 404

        course = db.query(Course).filter_by(id=student.course_id).first()
        if not course or course.department_id != department_id:
            flash("That student is not in your department.", "danger")
            return redirect(url_for("hod_dashboard"))

        sem = db.query(Semester).filter_by(id=semester_id, course_id=student.course_id).first()
        if not sem:
            return "Semester not found for this student", 404

        atype = get_course_assessment_type(course)
        subjects = (
            subject_query(db).filter_by(course_id=student.course_id, semester_id=sem.id)
            .order_by(Subject.name.asc()).all()
        )

        # ── Only include subjects the student actually has marks for ──
        rows = []
        grand_total = 0
        counted = 0
        for subj in subjects:
            info = get_student_subject_marks(db, student.id, subj.id, sem.id, course)
            if not info:
                continue
            rows.append({"subject": subj, "info": info})
            grand_total += info["total"]
            counted += 1
        avg = round(grand_total / counted, 2) if counted else None

        if not rows:
            flash(
                "No marks have been entered for this student in the selected semester yet.",
                "warning",
            )
            return redirect(url_for("hod_students_list"))

        labels = list(KASNEB_LABELS) if atype == "kasneb" else get_periodic_labels(course)

        return render_template(
            "students/print_report.html",
            student=student, semester=sem, rows=rows,
            assessment_type=atype,
            periodic_labels=labels,
            grand_total=grand_total, counted=counted, avg=avg,
            current_year=datetime.now().year, admin_mode=True,
        )
    finally:
        db.close()


@app.route("/hod/students/export")
def hod_students_export():
    if session.get("role") != "hod":
        return redirect(url_for("login"))

    department_id = session["department_id"]
    db = SessionLocal()
    try:
        department = db.query(Department).filter_by(id=department_id).first()
        if not department:
            flash("Department not found.", "danger")
            return redirect(url_for("logout"))

        courses = db.query(Course).filter_by(department_id=department_id).order_by(Course.name.asc()).all()
        course_ids = [c.id for c in courses]

        course_id = request.args.get("course_id", type=int)
        exam_type = request.args.get("exam_type", "").strip()
        search = request.args.get("search", "").strip()

        query = student_query(db)
        if course_ids:
            query = query.filter(StudentProfile.course_id.in_(course_ids))
        else:
            query = query.filter(False)

        if course_id and course_id in course_ids:
            query = query.filter(StudentProfile.course_id == course_id)
        if exam_type:
            query = query.filter(StudentProfile.exam_type == exam_type)
        if search:
            like = f"%{search}%"
            query = query.filter(or_(
                StudentProfile.full_name.ilike(like),
                StudentProfile.admission_number.ilike(like),
                StudentProfile.phone_number.ilike(like),
            ))

        students = query.order_by(StudentProfile.full_name.asc()).all()

        sel_course = next((c for c in courses if c.id == course_id), None)
        label_parts = [department.name.replace(" ", "_")]
        if sel_course:
            label_parts.append(sel_course.name.replace(" ", "_")[:30])
        if exam_type:
            label_parts.append(exam_type.upper())
        filename_label = "_".join(label_parts)

        data = []
        for s in students:
            data.append({
                "Admission No": s.admission_number or "",
                "Full Name": s.full_name or "",
                "Course": s.course.name if s.course else "",
                "Course Level": s.course.level if s.course else "",
                "Module": s.module or "",
                "Semester": (s.current_semester.name if s.current_semester else ""),
                "Exam Type": (s.exam_type or "").upper(),
                "Phone Number": s.phone_number or "",
                "Status": "Blocked" if s.blocked else "Active",
            })

        df = pd.DataFrame(data)
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Students")
            ws = writer.sheets["Students"]
            for idx, col in enumerate(df.columns, start=1):
                max_len = max(
                    [len(str(col))] + [len(str(v)) for v in df[col].tolist()]
                ) if len(df) else len(str(col))
                ws.column_dimensions[
                    ws.cell(row=1, column=idx).column_letter
                ].width = min(max(max_len + 2, 12), 40)

        buf.seek(0)

        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"Students_{filename_label}.xlsx",
        )
    finally:
        db.close()


"""
============================================================
STUDENT ROUTES
============================================================
"""
@app.route("/complete_profile", methods=["GET", "POST"])
def complete_profile():
    if "user_id" not in session:
        return redirect(url_for("login"))

    db = SessionLocal()
    user_id = session["user_id"]

    if request.method == "POST":
        full_name = request.form.get("full_name").strip()
        exam_type = request.form.get("exam_type")
        course_id = request.form.get("course_id")
        admission_number = request.form.get("admission_number").strip()
        phone_number = request.form.get("phone_number").strip()

        try:
            if db.query(StudentProfile).filter_by(admission_number=admission_number).first():
                flash("❌ Admission number already exists.", "danger")
                courses = db.query(Course).all()
                return render_template(
                    "students/complete_profile.html", courses=courses,
                    full_name=full_name, exam_type=exam_type,
                    admission_number=admission_number, phone_number=phone_number,
                )

            if db.query(StudentProfile).filter_by(user_id=user_id).first():
                flash("❌ You have already completed your profile.", "warning")
                return redirect(url_for("student_dashboard"))

            db.add(StudentProfile(
                full_name=full_name, exam_type=exam_type,
                course_id=course_id, admission_number=admission_number,
                phone_number=phone_number, user_id=user_id,
            ))
            db.commit()
            flash("✅ Profile completed successfully!", "success")
            return redirect(url_for("student_dashboard"))
        except Exception:
            db.rollback()
            flash("❌ Something went wrong. Try again.", "danger")
            courses = db.query(Course).all()
            return render_template("students/complete_profile.html", courses=courses)
        finally:
            db.close()

    courses = db.query(Course).all()
    db.close()
    return render_template("students/complete_profile.html", courses=courses)


@app.route("/student/dashboard")
def student_dashboard():
    if "username" not in session or session.get("role") != "student":
        flash("Please log in as a student first.", "error")
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        user = db.query(User).filter_by(username=session["username"]).first()
        if not user:
            flash("User not found.", "error")
            return redirect(url_for("logout"))

        student_profile = db.query(StudentProfile).filter_by(user_id=user.id).first()
        if not student_profile:
            flash("Complete your profile before proceeding.", "warning")
            return redirect(url_for("complete_profile"))

        course = student_profile.course
        atype = get_course_assessment_type(course)

        semesters = (
            semester_query(db).filter_by(course_id=student_profile.course_id)
            .order_by(Semester.created_at.asc()).all()
        )

        all_quizzes = (
            quiz_query(db).filter_by(course_id=student_profile.course_id, status="active")
            .order_by(Quiz.id.asc()).all()
        )
        quizzes_by_subject = {}
        for q in all_quizzes:
            quizzes_by_subject.setdefault(q.subject_id, []).append(q)

        results = db.query(Result).filter_by(student_id=student_profile.id).all()
        results_by_quiz = {r.quiz_id: r for r in results}
        taken_quiz_ids = set(results_by_quiz.keys())

        def build(subject):
            qs = quizzes_by_subject.get(subject.id, [])
            taken = sum(1 for q in qs if q.id in taken_quiz_ids)
            row = {
                "subject": subject, "quizzes": qs,
                "total": len(qs), "taken": taken,
                "percent": round(taken / len(qs) * 100) if qs else 0,
                "info": None,
            }
            if subject.semester_id:
                row["info"] = get_student_subject_marks(
                    db, student_profile.id, subject.id, subject.semester_id, course
                )
            return row

        all_subjects = (
            subject_query(db).filter_by(course_id=student_profile.course_id)
            .order_by(Subject.name.asc()).all()
        )
        all_subjects_data = [build(s) for s in all_subjects]

        current_semester = None
        semester_groups = []
        for sem in semesters:
            subs = (
                subject_query(db).filter_by(course_id=student_profile.course_id, semester_id=sem.id)
                .order_by(Subject.name.asc()).all()
            )
            if sem.is_current:
                current_semester = sem
            semester_groups.append({
                "semester": sem, "subjects": [build(s) for s in subs],
                "count": len(subs), "is_current": bool(sem.is_current),
            })

        unassigned = (
            subject_query(db).filter(Subject.course_id == student_profile.course_id)
            .filter(Subject.semester_id.is_(None))
            .order_by(Subject.name.asc()).all()
        )
        unassigned_subjects = [build(s) for s in unassigned]

        messages = (
            db.query(Message)
            .filter(
                (Message.target_type == "all") |
                ((Message.target_type == "course") & (Message.course_id == student_profile.course_id))
            )
            .order_by(Message.created_at.desc()).all()
        )

        return render_template(
            "students/student_dashboard.html",
            username=user.username, student_profile=student_profile,
            is_knec=(atype == "knec"),
            assessment_type=atype,
            semester_groups=semester_groups,
            current_semester=current_semester,
            unassigned_subjects=unassigned_subjects,
            all_subjects_data=all_subjects_data,
            total_subjects=len(all_subjects),
            taken_quiz_ids=list(taken_quiz_ids),
            results_by_quiz=results_by_quiz,
            messages_from_admin=messages,
            current_year=datetime.now().year,
        )
    finally:
        db.close()


@app.route("/take_exam/<int:quiz_id>", methods=["GET", "POST"])
def take_exam(quiz_id):
    if "user_id" not in session:
        flash("Please log in first.", "warning")
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        student = db.query(StudentProfile).filter_by(user_id=session["user_id"]).first()
        if not student:
            flash("Student profile not found.", "danger")
            return redirect(url_for("student_dashboard"))

        if student.blocked:
            flash("You can't access this exam. Please clear the fee to regain access.", "danger")
            return redirect(url_for("student_dashboard"))

        quiz = quiz_query(db).filter_by(id=quiz_id).first()
        if not quiz:
            flash("Quiz not found.", "danger")
            return redirect(url_for("student_dashboard"))

        db.query(ActivityLog).filter_by(student_id=student.id, is_active=True).update({"is_active": False})
        db.add(ActivityLog(student_id=student.id, activity_type="exam", is_active=True))
        db.commit()

        if request.method == "POST":
            score = 0
            total_marks = 0
            questions = db.query(Question).filter_by(quiz_id=quiz.id).order_by(Question.id.asc()).all()

            for q in questions:
                selected = request.form.get(f"question_{q.id}")
                total_marks += q.marks
                if selected and selected.strip().lower() == (q.correct_option or "").strip().lower():
                    score += q.marks

            percentage = (score / total_marks) * 100 if total_marks > 0 else 0

            db.add(Result(student_id=student.id, quiz_id=quiz.id,
                          score=score, total_marks=total_marks, percentage=percentage))

            db.query(ActivityLog).filter_by(
                student_id=student.id, activity_type="exam", is_active=True,
            ).update({"is_active": False})
            db.commit()

            return render_template(
                "students/take_exam.html",
                quiz=quiz, questions=[], student=student,
                show_result=True, score=score, total_marks=total_marks, percentage=percentage,
            )

        questions = db.query(Question).filter_by(quiz_id=quiz.id).order_by(Question.id.asc()).all()
        return render_template(
            "students/take_exam.html",
            quiz=quiz, questions=questions, student=student, show_result=False,
        )
    finally:
        db.close()


@app.route("/student/results")
def student_results():
    if "username" not in session or session.get("role") != "student":
        flash("Please log in as a student first.", "error")
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        user = db.query(User).filter_by(username=session["username"]).first()
        if not user:
            flash("User not found.", "error")
            return redirect(url_for("logout"))

        student = db.query(StudentProfile).filter_by(user_id=user.id).first()
        if not student:
            flash("Student profile not found.", "error")
            return redirect(url_for("logout"))

        course = student.course
        atype = get_course_assessment_type(course)

        if atype in ("short", "kasneb"):
            results = (
                db.query(Result)
                .options(
                    joinedload(Result.quiz).joinedload(Quiz.subject),
                    joinedload(Result.quiz).joinedload(Quiz.course),
                )
                .join(Quiz, Result.quiz_id == Quiz.id)
                .filter(Result.student_id == student.id)
                .order_by(Result.taken_on.desc()).all()
            )
            return render_template(
                "students/results.html",
                is_knec=False, results=results, student=student,
                assessment_type=atype,
                current_year=datetime.now().year,
            )

        semesters = (
            semester_query(db).filter_by(course_id=student.course_id)
            .order_by(Semester.created_at.asc()).all()
        )

        semester_reports = []
        for sem in semesters:
            subjects = (
                subject_query(db).filter_by(course_id=student.course_id, semester_id=sem.id)
                .order_by(Subject.name.asc()).all()
            )
            rows = []
            for subj in subjects:
                info = get_student_subject_marks(db, student.id, subj.id, sem.id, course)
                rows.append({"subject": subj, "info": info})
            semester_reports.append({
                "semester": sem, "rows": rows, "is_current": bool(sem.is_current),
            })

        labels = list(KASNEB_LABELS) if atype == "kasneb" else get_periodic_labels(course)

        return render_template(
            "students/results.html",
            is_knec=(atype == "knec"), student=student,
            semester_reports=semester_reports,
            assessment_type=atype,
            periodic_labels=labels,
            current_year=datetime.now().year,
        )
    finally:
        db.close()

@app.route("/student/print_report")
def student_print_report():
    if "username" not in session or session.get("role") != "student":
        flash("Please log in as a student first.", "error")
        return redirect(url_for("login"))

    db = SessionLocal()
    try:
        user = db.query(User).filter_by(username=session["username"]).first()
        if not user:
            return redirect(url_for("logout"))

        student = db.query(StudentProfile).filter_by(user_id=user.id).first()
        if not student:
            return redirect(url_for("logout"))

        if (student.exam_type or "").strip().lower() != "knec":
            flash("Printable reports are only available for KNEC students.", "warning")
            return redirect(url_for("student_results"))

        course = student.course
        atype = get_course_assessment_type(course)

        sem_id = request.args.get("semester_id", type=int)
        if sem_id:
            sem = db.query(Semester).filter_by(id=sem_id, course_id=student.course_id).first()
        else:
            sem = db.query(Semester).filter_by(course_id=student.course_id, is_current=True).first()

        if not sem:
            flash("No semester selected.", "warning")
            return redirect(url_for("student_results"))

        subjects = (
            subject_query(db).filter_by(course_id=student.course_id, semester_id=sem.id)
            .order_by(Subject.name.asc()).all()
        )

        # ── Only include subjects the student actually has marks for ──
        rows = []
        grand_total = 0
        counted = 0
        for subj in subjects:
            info = get_student_subject_marks(db, student.id, subj.id, sem.id, course)
            if not info:
                continue
            rows.append({"subject": subj, "info": info})
            grand_total += info["total"]
            counted += 1

        avg = round(grand_total / counted, 2) if counted else None

        if not rows:
            flash(
                "No marks have been entered for this student in the selected semester yet.",
                "warning",
            )
            return redirect(url_for("student_results"))

        labels = list(KASNEB_LABELS) if atype == "kasneb" else get_periodic_labels(course)

        return render_template(
            "students/print_report.html",
            student=student, semester=sem, rows=rows,
            assessment_type=atype,
            periodic_labels=labels,
            grand_total=grand_total, counted=counted, avg=avg,
            current_year=datetime.now().year,
        )
    finally:
        db.close()



"""
============================================================
ENTRY POINT
============================================================
"""
if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)